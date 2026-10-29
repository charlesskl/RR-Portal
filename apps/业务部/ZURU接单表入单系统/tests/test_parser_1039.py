"""Synthetic PO review: real XLSX packages and generated schedule readback only."""
import hashlib
import io
import os
from pathlib import Path
import tempfile
import unittest
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape
import openpyxl
import po_parser

STATE=tempfile.TemporaryDirectory(prefix='pr1039-parser-app-')
os.environ['DATA_PATH']=STATE.name
os.environ['LOCAL_MODE']='0'
import app as service

def xlsx(path, spec='S001-FUGGLER-SYNTHETIC', sku='77485GQ3', cells=(), shapes=(), prefix='a'):
    book=openpyxl.Workbook(); sheet=book.active
    for n,v in {'A6':'Date:', 'B6':'2026.10.20', 'A7':'PO #:', 'G7':'4500000001', 'A8':'Customer Name:', 'G8':'Synthetic Customer', 'U5':'Destination Country:', 'V5':'USA'}.items():sheet[n]=v
    for col,v in enumerate(['Line','SKU','SKU-Spec','Name','Price USD','Qty','Delivery Date','Customer PO'],1):sheet.cell(14,col,v)
    for col,v in enumerate([1,sku,spec,'Synthetic product',1.25,10,'2026.11.01','SYNTH-US'],1):sheet.cell(15,col,v)
    sheet['A16']='Total'
    for row,value in enumerate(cells,20):sheet.cell(row,1,value)
    book.save(path);book.close()
    if shapes:
        xml='<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:'+prefix+'="http://schemas.openxmlformats.org/drawingml/2006/main">'
        for runs in shapes:xml+='<xdr:twoCellAnchor><xdr:sp><xdr:txBody><'+prefix+':p>'+''.join('<'+prefix+':r><'+prefix+':t>'+escape(t)+'</'+prefix+':t></'+prefix+':r>'for t in runs)+'</'+prefix+':p></xdr:txBody></xdr:sp></xdr:twoCellAnchor>'
        xml+='</xdr:wsDr>'
        with zipfile.ZipFile(path,'a')as z:z.writestr('xl/drawings/review.xml',xml)

class Parser1039(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='pr1039-po-');self.addCleanup(self.tmp.cleanup);self.path=Path(self.tmp.name,'PO-4500000001.xlsx')
    def test_ordinary_notes_and_cross_shape_fragments_never_cancel(self):
        for note in ('Cancel Date','avoid shipping damage','Line 10 cancelled','订单取消日期','此前已取消的订单说明'):
            with self.subTest(note=note):
                xlsx(self.path,cells=[note],shapes=[(note,)]);self.assertFalse(po_parser.is_cancelled(self.path))
        xlsx(self.path,shapes=[('CAN',),('CEL',)]);self.assertFalse(po_parser.is_cancelled(self.path))
    def test_exact_cell_stamps_and_independent_rich_text_shapes_cancel(self):
        for stamp in ('CANCEL','CANCEL!','订单取消！','取消','VOID','Cancelled!'):
            with self.subTest(cell=stamp):xlsx(self.path,cells=[stamp]);self.assertTrue(po_parser.is_cancelled(self.path))
        xlsx(self.path,shapes=[('Cancel Date',),('订','单','取','消')]);self.assertTrue(po_parser.is_cancelled(self.path))
    def test_equivalent_drawing_namespace_prefix_does_not_hide_stamp(self):
        xlsx(self.path,shapes=[('CANCEL',)],prefix='text');self.assertTrue(po_parser.is_cancelled(self.path))
    def test_sku_fallback_retains_spec_priority_and_all_parsed_fields(self):
        for spec,sku,expected in [('S001-FUGGLER-SYNTHETIC','77485GQ3','77485'),('', 'MEC421-S001','MEC421'),('9574-STD','77485GQ3','9574'),('INVALID-SPEC','INVALID-SKU','')]:
            with self.subTest(spec=spec):
                xlsx(self.path,spec,sku);before=self.path.read_bytes();p=po_parser.parse(self.path);v=p['lines'][0]
                self.assertEqual(v['simple_no'],expected);self.assertEqual(v['sku'],sku);self.assertEqual(v['sku_spec'],spec)
                self.assertEqual((v['qty'],v['price'],v['delivery']),(10,1.25,datetime(2026,11,1)));self.assertEqual(self.path.read_bytes(),before)
    def test_real_upload_and_export_keeps_ordinary_po_and_filters_exact_cancel(self):
        service.app.config['TESTING']=True
        for key,name in [('UPLOAD_FOLDER','uploads'),('EXPORT_FOLDER','exports'),('DATA_FOLDER','data')]:
            folder=Path(self.tmp.name,name);folder.mkdir();service.app.config[key]=str(folder)
        service.ITEM_MAP={'77485':{'cn_name':'合成品名','workshop':'合成车间','schedule_name':'合成排期'}};service.CANCELLED_POS=set();client=service.app.test_client()
        xlsx(self.path,shapes=[('Cancel Date',),('avoid shipping damage',)]);before=self.path.read_bytes()
        r=client.post('/api/process',data={'files':(io.BytesIO(before),'PO-4500000001.xlsx')});self.assertEqual(r.status_code,200);v=r.get_json();self.assertEqual(v['written'],1);self.assertEqual(v.get('cancelled',[]),[])
        exports=list(Path(service.app.config['EXPORT_FOLDER']).glob('*.xlsx'));self.assertEqual(len(exports),1);book=openpyxl.load_workbook(exports[0],data_only=False)
        try:
            self.assertEqual(book['排期']['C3'].value,77485);self.assertEqual(book['排期']['D3'].value,'77485GQ3');self.assertTrue(any('SKU-SPEC' in warning for warning in v['warnings']));self.assertEqual(book['排期']['F3'].value,10);self.assertEqual(book['排期']['G3'].value,1.25);self.assertEqual(book['排期']['H3'].value,'=F3*G3');self.assertEqual(book['排期']['I3'].value,'=H3*7.75');self.assertEqual(book['排期']['K3'].value,'美国')
        finally:book.close()
        self.assertEqual(self.path.read_bytes(),before);xlsx(self.path,cells=['CANCEL']);r=client.post('/api/process',data={'files':(io.BytesIO(self.path.read_bytes()),'PO-4500000001.xlsx')});self.assertEqual(r.status_code,200);self.assertEqual(r.get_json()['written'],0);self.assertEqual(r.get_json()['cancelled'],['PO-4500000001.xlsx'])

def tearDownModule():
    for handler in list(service.logging.getLogger().handlers):
        if handler is service._log_handler:service.logging.getLogger().removeHandler(handler);handler.close()
    STATE.cleanup()
if __name__=='__main__':unittest.main()
