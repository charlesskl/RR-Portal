"""Synthetic real Flask/XLSX regressions; no network or business filesystem."""
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
ISOLATED = tempfile.TemporaryDirectory(prefix='zuru1039-app-')
os.environ['DATA_PATH'] = ISOLATED.name
os.environ['LOCAL_MODE'] = '0'
sys.path.insert(0, str(ROOT))
import app as service

def workbook_bytes(book):
    data = io.BytesIO(); book.save(data); book.close(); return data.getvalue()

def schedule(pos=('4500000001',), title='排期', other_active=False, duplicate=False):
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = title
    sheet['B2'] = '合同'
    for i, value in enumerate(pos, 3): sheet.cell(i, 2, value)
    if other_active:
        other = book.create_sheet('货号'); other['B3'] = 'NOT-A-PO'; book.active = 1
    if duplicate: book.create_sheet(' 排期 ')['B3'] = 'DIFFERENT'
    return workbook_bytes(book)

def po(specs=('111-S001',), number='4500000001'):
    book = openpyxl.Workbook(); sheet = book.active
    for address,value in {'A6':'Date:', 'B6':'2026.10.20', 'A7':'PO #:', 'G7':number,
                          'A8':'Customer Name:', 'G8':'Synthetic Customer',
                          'U5':'Destination Country:', 'V5':'Hong Kong'}.items(): sheet[address] = value
    for col,value in enumerate(['Line','SKU','SKU-Spec','Name','Price USD','Qty','Delivery Date','Customer PO'],1): sheet.cell(14,col,value)
    for row,spec in enumerate(specs,15):
        for col,value in enumerate([row-14,'111',spec,'Synthetic toy',1.25,10,'2026.11.01','TEST-HK'],1): sheet.cell(row,col,value)
    sheet.cell(15+len(specs),1,'Total')
    return workbook_bytes(book)

class App1039(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=ISOLATED.name); self.addCleanup(self.temp.cleanup)
        for key,name in [('UPLOAD_FOLDER','uploads'),('EXPORT_FOLDER','exports'),('DATA_FOLDER','data')]:
            p=Path(self.temp.name,name);p.mkdir();service.app.config[key]=str(p)
        service.app.config['TESTING']=True
        self.cfg=Path(service.app.config['DATA_FOLDER'],'config.json')
        for name,value in [('_CONFIG_FILE',str(self.cfg)),('_pos_cache',{'key':None,'pos':set()}),
                           ('ITEM_MAP',{'111':{'cn_name':'合成品','workshop':'A','schedule_name':'Synthetic'}}),
                           ('CANCELLED_POS',set()),('_tasks',{})]:
            change=patch.object(service,name,value);change.start();self.addCleanup(change.stop)
        self.client=service.app.test_client()

    def configured(self, body):
        table=Path(self.temp.name,'selected.xlsx');table.write_bytes(body)
        self.cfg.write_text(json.dumps({'order_table_path':str(table),'keep':'original'}))
        return table

    def process(self, body=None, name='PO-4500000001.xlsx'):
        return self.client.post('/api/process',data={'files':(io.BytesIO(body or po()),name)})

    def test_named_schedule_not_active_and_only_trailing_numeric_decimal_removed(self):
        table=self.configured(schedule((4500000001.0,'450.012.0',' 4500000002.0 '),title=' 排期 ',other_active=True))
        self.assertEqual(service._read_existing_pos(str(table)),{'4500000001','450.012','4500000002'})

    def test_real_upload_classify_and_process_use_schedule_and_never_wrong_active(self):
        response=self.client.post('/api/upload_table',data={'file':(io.BytesIO(schedule(other_active=True)),'synthetic.xlsx')})
        self.assertEqual(response.status_code,200);self.assertEqual(response.json['count'],1)
        response=self.client.post('/api/classify',data={'files':(io.BytesIO(po()),'PO-4500000001.xlsx')})
        self.assertEqual(response.status_code,200)
        rows=response.json['results'] if isinstance(response.json,dict) else response.json
        self.assertTrue(rows[0]['is_revision'])
        processed=self.process();self.assertEqual(processed.status_code,200)
        self.assertEqual(processed.json['written'],0);self.assertEqual(processed.json['revisions'],['PO-4500000001.xlsx'])
        self.assertEqual(list(Path(service.app.config['EXPORT_FOLDER']).iterdir()),[])

    def test_bad_upload_preserves_existing_valid_source_and_configuration(self):
        table=self.configured(schedule());before=table.read_bytes();cfg=self.cfg.read_bytes()
        files=set(Path(service.app.config['DATA_FOLDER']).iterdir())
        response=self.client.post('/api/upload_table',data={'file':(io.BytesIO(schedule(title='货号')),'wrong.xlsx')})
        self.assertEqual(response.status_code,400);self.assertIn('排期',response.json['error'])
        self.assertEqual(table.read_bytes(),before);self.assertEqual(self.cfg.read_bytes(),cfg)
        self.assertEqual(set(Path(service.app.config['DATA_FOLDER']).iterdir()),files)
        self.assertEqual(self.process().json['written'],0)

    def test_atomic_config_publish_failure_preserves_source_and_config(self):
        accepted=self.client.post('/api/upload_table',data={'file':(io.BytesIO(schedule()),'original.xlsx')})
        self.assertEqual(accepted.status_code,200)
        table=Path(json.loads(self.cfg.read_text())['order_table_path'])
        before=table.read_bytes();cfg=self.cfg.read_bytes()
        files=set(Path(service.app.config['DATA_FOLDER']).iterdir())
        with patch.object(service.os,'replace',side_effect=OSError('synthetic publication failure')):
            response=self.client.post('/api/upload_table',data={'file':(io.BytesIO(schedule(('4500000002',))),'new.xlsx')})
        self.assertEqual(response.status_code,500)
        self.assertEqual(table.read_bytes(),before);self.assertEqual(self.cfg.read_bytes(),cfg)
        self.assertEqual(set(Path(service.app.config['DATA_FOLDER']).iterdir()),files)
        self.assertEqual(self.process().json['written'],0)

    def test_corrupt_upload_preserves_previously_uploaded_workbook_config_and_comparison(self):
        accepted=self.client.post('/api/upload_table',data={'file':(io.BytesIO(schedule()),'original.xlsx')})
        self.assertEqual(accepted.status_code,200)
        table=Path(json.loads(self.cfg.read_text())['order_table_path'])
        before=table.read_bytes();cfg=self.cfg.read_bytes()
        files=set(Path(service.app.config['DATA_FOLDER']).iterdir())
        response=self.client.post('/api/upload_table',data={'file':(io.BytesIO(b'broken zip'),'broken.xlsx')})
        self.assertEqual(response.status_code,400)
        self.assertEqual(table.read_bytes(),before);self.assertEqual(self.cfg.read_bytes(),cfg)
        self.assertEqual(set(Path(service.app.config['DATA_FOLDER']).iterdir()),files)
        processed=self.process();self.assertEqual(processed.status_code,200)
        self.assertEqual(processed.json['written'],0)
        self.assertEqual(processed.json['revisions'],['PO-4500000001.xlsx'])

    def test_missing_schedule_process_fails_before_classification_and_cleans_upload(self):
        table=self.configured(schedule(title='货号'));before=table.read_bytes();cfg=self.cfg.read_bytes()
        response=self.process()
        self.assertEqual(response.status_code,400);self.assertIn('排期',response.json['error'])
        self.assertEqual(table.read_bytes(),before);self.assertEqual(self.cfg.read_bytes(),cfg)
        self.assertEqual(list(Path(service.app.config['UPLOAD_FOLDER']).iterdir()),[])
        self.assertEqual(list(Path(service.app.config['EXPORT_FOLDER']).iterdir()),[])
        self.assertEqual(service._tasks,{})

    def test_empty_schedule_classify_process_and_local_check_fail_explicitly(self):
        table=self.configured(schedule(()))
        for endpoint in ('/api/classify','/api/process'):
            response=self.client.post(endpoint,data={'files':(io.BytesIO(po()),'PO-4500000001.xlsx')})
            self.assertEqual(response.status_code,400,endpoint);self.assertIn('PO',response.json['error'])
        with patch.object(service,'LOCAL_MODE',True):
            response=self.client.post('/api/check_path',json={'path':str(table)})
        self.assertEqual(response.status_code,400)
        self.assertEqual(list(Path(service.app.config['UPLOAD_FOLDER']).iterdir()),[])
        self.assertEqual(list(Path(service.app.config['EXPORT_FOLDER']).iterdir()),[])

    def test_configured_missing_or_corrupt_table_never_uses_filename_fallback(self):
        table=self.configured(b'not an Excel zip')
        for remove in (False,True):
            if remove:table.unlink()
            response=self.process()
            self.assertEqual(response.status_code,400)
            self.assertEqual(list(Path(service.app.config['UPLOAD_FOLDER']).iterdir()),[])
            self.assertEqual(list(Path(service.app.config['EXPORT_FOLDER']).iterdir()),[])

    def test_ambiguous_normalized_schedule_fails_instead_of_choosing_arbitrary_data(self):
        table=self.configured(schedule(duplicate=True))
        with self.assertRaisesRegex(ValueError,'多个'):service._read_existing_pos(str(table))

    def test_sku_spec_fallback_real_api_and_download_workbook_preserve_amount_country(self):
        response=self.process(po(('S001-FUGGLER-SYNTHETIC-NAME',)))
        self.assertEqual(response.status_code,200);self.assertEqual(response.json['written'],1)
        self.assertTrue(any('SKU-SPEC' in value for value in response.json['warnings']))
        download=self.client.get('/api/download/'+response.json['filename']);self.assertEqual(download.status_code,200)
        book=openpyxl.load_workbook(io.BytesIO(download.data),data_only=False)
        try:
            sheet=book['排期']
            self.assertEqual(sheet['C3'].value,111);self.assertEqual(sheet['D3'].value,'111')
            self.assertEqual(sheet['F3'].value,10);self.assertEqual(sheet['G3'].value,1.25)
            self.assertEqual(sheet['H3'].value,'=F3*G3');self.assertEqual(sheet['I3'].value,'=H3*7.75')
            self.assertEqual(sheet['K3'].value,'香港-亚洲');self.assertEqual(sheet.freeze_panes,'A3')
            self.assertEqual(book['货号']['A2'].value,111)
        finally:book.close();download.close()

    def test_existing_slb_plb_and_stdproduct_exports_still_use_their_suffix_rules(self):
        response=self.process(po(('111-SLB-S001','111-PLB-S001','111-STDPRODUCT')))
        self.assertEqual(response.status_code,200);self.assertEqual(response.json['written'],3)
        book=openpyxl.load_workbook(Path(service.app.config['EXPORT_FOLDER'],response.json['filename']))
        try:
            self.assertEqual([book['排期'].cell(i,4).value for i in (3,4,5)],['111SLB','111PLB','111'])
            self.assertEqual([book['排期'].cell(i,6).value for i in (3,4,5)],[10,10,10])
            self.assertEqual([book['排期'].cell(i,7).value for i in (3,4,5)],[1.25]*3)
        finally:book.close()

    def test_unconfigured_filename_revision_compatibility_remains(self):
        response=self.process(name='PO-4500000001-REV.xlsx')
        self.assertEqual(response.status_code,200);self.assertEqual(response.json['written'],0)
        self.assertEqual(response.json['revisions'],['PO-4500000001-REV.xlsx'])

def tearDownModule():
    # Only this module's import-time log file may be closed here.
    handler = service._log_handler
    filename = getattr(handler, 'baseFilename', None)
    if filename and Path(filename).resolve().is_relative_to(Path(ISOLATED.name).resolve()):
        root = service.logging.getLogger()
        if handler in root.handlers:
            root.removeHandler(handler)
        handler.close()
    ISOLATED.cleanup()

if __name__=='__main__':unittest.main()
