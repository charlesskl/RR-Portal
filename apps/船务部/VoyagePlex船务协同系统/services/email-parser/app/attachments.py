"""Attachment parsers adapted from RR-Portal main@690cbab.

This module owns its contracts and has no runtime dependency on the legacy tree.
"""

import re
from pathlib import Path
from typing import Optional

import openpyxl
import pdfplumber
from docx import Document


HEADER_ALIASES = {
    "product_code": ("sku", "item no", "item#", "product code", "货号", "keycode", "style"),
    "product_name": ("description", "product name", "item description", "货名", "品名"),
    # 业务“数量”是产品总数（Retail Unit）；旧模板没有该列时再回退到 Shipping Quantity。
    "quantity": ("retail unit", "shipping quantity", "order qty", "quantity", "qty", "数量"),
    "pieces": ("no of carton", "no carton", "no. of cartons", "total carton", "carton qty", "cartons", "ctns", "件数"),
    "spec": ("pc per carton", "pcs/ctn", "case pack", "inners", "每箱"),
    "volume": ("total cbm", "cbm", "volume", "measurement", "体积"),
    "customer_po": ("customer po", "order no", "客人po", "客户po"),
    "contract_number": ("zuru po", "contract", "订单号", "合同号", "po no"),
    "gross_weight": ("gross weight", "g.w", "gw", "毛重"),
    "net_weight": ("net weight", "n.w", "nw", "净重"),
    "supplier": ("actual factory", "supplier", "factory", "工厂"),
    "factory_remark": ("main assembly factory", "factory remark", "装配工厂"),
    "pallet_count": ("no of pallet", "no. of pallets", "pallet qty", "卡板数"),
    "cargo_receipt": ("cargo receipt", "receipt no", "落货纸"),
    "length": ("length", "长", "l"),
    "width": ("width", "宽", "w"),
    "height": ("height", "高", "h"),
}

PRODUCT_NAME_ALIASES = (
    ("ROBO ALIVE-DINOSAUR", "恐龙"),
    ("MAGIC BIRD", "魔法小鸟"),
    ("BRAINROT", "大脑"),
    ("FUGGLER", "非凡系列"),
)


def parse_attachment(path: Path, original_name: str) -> dict:
    suffix = Path(original_name).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return parse_excel(path)
    if suffix == ".xls":
        return parse_xls(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_word(path)
    return {"kind": "unsupported", "fields": {}, "items": [], "warnings": []}


def parse_excel(path: Path) -> dict:
    # read_only 模式流式读取单元格值，不加载样式/图形：大附件在 512MB 容器里
    # 用普通模式加载会把整个工作簿（含 DrawingML）读进内存，既慢又有 OOM 风险。
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        yax = _parse_yax(wb)
        if yax:
            return {"kind": "yax", "fields": yax, "items": yax.pop("po_so_mapping", []), "warnings": []}

        sheet = _choose_packing_sheet(wb)
        sheet_title = sheet.title
        rows = _read_sheet_rows(sheet)
    finally:
        wb.close()
    header_row, headers = _find_header(rows)
    columns = _column_map(headers)
    _apply_measurement_columns(headers, columns)
    items = []
    blank_streak = 0
    for row_index, values in enumerate(rows[header_row:], header_row + 1):
        product = _value(values, columns.get("product_code"))
        if not product or _is_non_cargo_label(product):
            blank_streak += 1
            if blank_streak >= 10 and items:
                break
            continue
        blank_streak = 0
        product_name, specification = _product_name_and_spec(
            _value(values, columns.get("product_name")), _value(values, columns.get("spec")))
        quantity = _number(_value(values, columns.get("quantity")))
        pieces = _number(_value(values, columns.get("pieces")))
        item = {
            "product_code": str(product).strip(),
            "product_name": product_name,
            "quantity": quantity,
            "pieces": _pieces(pieces, quantity, specification),
            "spec": specification,
            "volume": _number(_value(values, columns.get("volume"))),
            "customer_po": _text(_value(values, columns.get("customer_po"))),
            "contract_number": _text(_value(values, columns.get("contract_number"))),
            "gross_weight": _number(_value(values, columns.get("gross_weight"))),
            "net_weight": _number(_value(values, columns.get("net_weight"))),
            "supplier": _text(_value(values, columns.get("supplier"))),
            "factory_remark": _text(_value(values, columns.get("factory_remark"))),
            "pallet_count": _number(_value(values, columns.get("pallet_count"))),
            "cargo_receipt": _text(_value(values, columns.get("cargo_receipt"))),
            "source_row": row_index,
        }
        dims = [_number(_value(values, columns.get(key))) for key in ("length", "width", "height")]
        item["box_dimensions"] = "*".join(_display_number(x) for x in dims) if all(x is not None for x in dims) else ""
        items.append(item)
    warnings = [] if items else [f"工作表“{sheet_title}”未识别到Packing List明细"]
    fields = {"loading_factory": "兴信"} if _xingxin_is_loading_factory(rows) else {}
    return {"kind": "packing_list", "fields": fields, "items": items, "warnings": warnings}


def parse_xls(path: Path) -> dict:
    """Parse legacy Excel files without modifying or converting the source."""
    import xlrd
    wb = xlrd.open_workbook(path)
    sheet = wb.sheet_by_index(0)
    rows = [[sheet.cell_value(row, col) for col in range(sheet.ncols)] for row in range(sheet.nrows)]
    flat = "\n".join(" | ".join(_text(value) for value in row) for row in rows)

    # Container Loading Plan/出货明细表 is a form, not a Packing List.
    if "Container Loading Plan" in flat or "出貨明細表" in flat:
        fields = {
            "destination_port": _row_label_value(rows, ("Port of Discharge",)),
            "port": _normalize_loading_port(_row_label_value(rows, ("Port of Loading",))),
            "container_type": _normalize_container(_row_label_value(rows, ("Container Size",))),
            "container_number": _row_label_value(rows, ("Container Number",)),
            "seal_number": _row_label_value(rows, ("Seal Number",)),
        }
        return {"kind": "container_loading_plan", "fields": fields, "items": [], "warnings": []}

    header_index, headers = max(
        ((index, row) for index, row in enumerate(rows[:40])),
        key=lambda candidate: len(_column_map(candidate[1])),
        default=(0, rows[0] if rows else []),
    )
    columns = _column_map(headers)
    items = []
    for row_index, values in enumerate(rows[header_index + 1:], header_index + 2):
        product = _value(values, columns.get("product_code"))
        if not product or _is_non_cargo_label(product):
            continue
        product_name, specification = _product_name_and_spec(
            _value(values, columns.get("product_name")), _value(values, columns.get("spec")))
        quantity = _number(_value(values, columns.get("quantity")))
        pieces = _number(_value(values, columns.get("pieces")))
        items.append({
            "product_code": _text(product),
            "product_name": product_name,
            "quantity": quantity,
            "pieces": _pieces(pieces, quantity, specification),
            "spec": specification,
            "volume": _number(_value(values, columns.get("volume"))),
            "customer_po": _text(_value(values, columns.get("customer_po"))),
            "contract_number": _text(_value(values, columns.get("contract_number"))),
            "source_row": row_index,
        })
    return {"kind": "packing_list_xls", "fields": {}, "items": items,
            "warnings": [] if items else ["旧版XLS未识别到货物明细，已保留表单字段供人工确认"]}


def parse_pdf(path: Path) -> dict:
    pages = []
    with pdfplumber.open(path) as document:
        for page in document.pages[:8]:
            pages.append(page.extract_text() or "")
    text = "\n".join(pages)
    delivery_address = _clean_address(_first(text, [
            r"PLACE\s*OF\s*DELIVERY\s*(?:\([^)]*\))?\s*[：:]\s*(.+?)(?:\n|$)",
            r"Final\s*Destination\s*(?:\([^)]*\))?\s*[：:]\s*(.+?)(?:\n|$)",
            r"Port\s*of\s*Discharge\s*(?:\([^)]*\))?\s*[：:]\s*(.+?)(?:\n|$)",
            r"(?:送货地址|交货地址|收货地址)[：:\s]*(.+?)(?:\n|$)",
        ]))
    fields = {
        "delivery_address": delivery_address,
        "customs_cutoff": _near_date(text, r"(?:Port\s*Cargo\s*Cut[\s-]*Off|VGM\s*Cut[\s-]*Off|CY\s*Closing|截重柜时间|截关时间)"),
        "si_deadline": _near_date(text, r"(?:Shipping\s*Instruction|SI\s*Cut[\s-]*Off|截补料时间)"),
        "port": _port(text),
        "country": _country(text),
    }
    return {"kind": "booking_pdf", "fields": fields, "items": [], "warnings": [] if text else ["PDF未提取到文本，可能是扫描件"]}


def parse_word(path: Path) -> dict:
    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    return {
        "kind": "word",
        "fields": {
            "warehouse": _first(text, [r"(?:仓库地址|送货地址|收货地址)[：:\s]*(.+?)(?:\n|$)"]),
            "so_number": _first(text, [r"SO\s*#?\s*([A-Z0-9]+)"]),
        },
        "items": [], "warnings": [],
    }


def _parse_yax(wb) -> Optional[dict]:
    values = []
    for ws in wb.worksheets[:3]:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 80, 80), values_only=True):
            values.append(["" if value is None else str(value).strip() for value in row])
    flat = "\n".join(" | ".join(row) for row in values)
    if "YAX" not in flat.upper() and "并柜拖车通知单" not in flat:
        return None
    fields = {
        "so_number": _first(flat, [r"\b(YAX\d{5,})\b"]),
        "booking_number": _value_after_label(values, ("船公司订舱单号", "订舱单号")),
        "container_type": _normalize_container(_value_after_label(values, ("柜型",))),
        "customs_cutoff": _value_after_label(values, ("截关时间", "截重柜时间")),
        "si_deadline": _value_after_label(values, ("截补料时间", "SI截止")),
        "po_so_mapping": [],
    }
    for row_index, row in enumerate(values):
        for col_index, cell in enumerate(row):
            if re.fullmatch(r"YAX\d{5,}", cell, re.I):
                po = next((value for value in row[:col_index] if re.fullmatch(r"\d{8,}", value)), "")
                cbm = next((_number(value) for value in row[col_index + 1:] if _number(value) is not None), None)
                fields["po_so_mapping"].append({"customer_po": po, "so_number": cell, "volume": cbm, "source_row": row_index + 1})
    return fields


def _choose_packing_sheet(wb):
    candidates = [ws for ws in wb.worksheets if any(x in ws.title.lower() for x in ("packing", "分柜", "明细", "detail"))]
    return candidates[0] if candidates else wb.active


def _read_sheet_rows(sheet) -> list[list]:
    """Stream sheet values into a plain list, stopping early on a long run of empty
    rows so an inflated sheet dimension cannot blow up memory or parse time."""
    # 部分非 Excel 工具生成的文件 dimension 标记偏小，会截断真实数据，按实际内容流式读取
    sheet.reset_dimensions()
    rows = []
    blank_streak = 0
    for row in sheet.iter_rows(values_only=True):
        values = list(row)
        if any(value is not None for value in values):
            blank_streak = 0
        else:
            blank_streak += 1
            if blank_streak >= 50:
                break
        rows.append(values)
    return rows


def _find_header(rows: list[list]) -> tuple[int, list]:
    best = (1, rows[0] if rows else [], 0)
    for index, headers in enumerate(rows[:40]):
        score = len(_column_map(headers))
        if score > best[2]:
            best = (index + 1, headers, score)
    return best[0], best[1]


def _column_map(headers: list) -> dict:
    result = {}
    normalized_headers = [re.sub(r"\s+", " ", str(raw or "").replace("\n", " ").strip().lower()) for raw in headers]
    for field, aliases in HEADER_ALIASES.items():
        candidates = []
        for index, normalized in enumerate(normalized_headers):
            for priority, alias in enumerate(aliases):
                if normalized == alias or (len(alias) > 2 and alias in normalized):
                    candidates.append((priority, 0 if normalized == alias else 1, index))
        if candidates:
            result[field] = min(candidates)[2]
    return result


def _apply_measurement_columns(headers: list, columns: dict) -> None:
    """Handle PL templates where L×W×H and total CBM share one merged heading."""
    normalized = [re.sub(r"\s+", " ", str(value or "").replace("\n", " ").strip().lower()) for value in headers]
    for index, header in enumerate(normalized):
        if ("measm" in header or "measurement" in header) and index + 5 < len(headers):
            columns.update({"length": index, "width": index + 2, "height": index + 4, "volume": index + 5})
            return


def _is_non_cargo_label(value) -> bool:
    normalized = re.sub(r"[\s:：_-]+", "", _text(value)).upper()
    return normalized in {
        "TOTAL", "SUBTOTAL", "GRANDTOTAL", "合计", "总计", "汇总",
        "SKUNO.", "SKUNO", "SKU", "货号",
    }


def _xingxin_is_loading_factory(rows: list[list]) -> bool:
    """Detect the customer-designated loading factory without exposing it as cargo ownership."""
    for row in rows:
        for value in row:
            text = _text(value)
            if re.search(r"(?:兴信|新信|HANSON).{0,12}(?:装柜|做柜|装货|LOADING)", text, re.I):
                return True
    return False


def _value(values, index):
    return values[index] if index is not None and index < len(values) else None


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value):
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        match = re.search(r"-?\d+(?:\.\d+)?", str(value))
        return float(match.group()) if match else None


def _product_name_and_spec(product_name, explicit_spec) -> tuple[str, Optional[float]]:
    """Keep product name and carton specification as separate business fields."""
    name = _text(product_name)
    pattern = r"(\d+(?:\.\d+)?)\s*(?:个\s*/?\s*箱|PCS?\s*/\s*(?:(?:PDQ|ENDCAP)\s*/\s*)?CTN|PCS?\s*/\s*ENDCAP|PCS?\s+PER\s+CARTON)\b"
    # The DESCRIPTION is the business-facing packing specification. Some templates put
    # display-box count (for example 1 ENDCAP/CTN) in PC PER CARTON, so it cannot override
    # an explicit 32PCS/ENDCAP or 41PCS/PDQ/CTN description.
    match = re.search(pattern, name, re.I) if name else None
    described_spec = _number(match.group(1)) if match else None
    column_spec = _number(explicit_spec)
    # Broken line wrapping in source workbooks can turn 13PCS into "1 3PCS". Taking the
    # larger valid value also preserves 32PCS/ENDCAP over the display-count value 1.
    specification = max(value for value in (described_spec, column_spec) if value is not None) \
        if described_spec is not None or column_spec is not None else None
    if specification is not None and name:
        name = re.sub(r"[,，;；\s]*" + pattern, "", name, flags=re.I).strip(" ,，;；")
    return _business_product_name(name), specification


def _business_product_name(source: str) -> str:
    """Extract the product identity from DESCRIPTION without treating packing text as its name."""
    name = re.sub(r"\s+", " ", source.replace("_x000D_", " ")).strip(" -/,，")
    name = re.sub(r"\bM\s+IXED\b", "MIXED", name, flags=re.I)
    name = re.sub(r"\bSURP\s+RISE\b", "SURPRISE", name, flags=re.I)
    if not name:
        return ""

    upper_name = name.upper()
    matches = [label for keyword, label in PRODUCT_NAME_ALIASES if keyword in upper_name]
    # A mixed carton can contain several independently named products. In that case the
    # confirmed business names are clearer than the very long packing description.
    if len(matches) > 1:
        return " / ".join(dict.fromkeys(matches))

    # DESCRIPTION usually starts with internal range/market tags, followed by the actual
    # product identity, and ends with display/carton packaging instructions.
    name = re.sub(r"^(?:(?:S\d{3}|EUR|MTS|NB|US|UK|STD)\s*[-/]\s*)+", "", name, flags=re.I)
    name = re.split(r"\s+MIXED\s+IN\s+(?:CTN|PDQ|ENDCAP)\b", name, maxsplit=1, flags=re.I)[0]
    name = re.split(
        r"[-,]\s*(?:PLUSH|CAPSULE|BLIND\s+BOX|COLOR\s+BOX|WINDOW\s+BOX|SHRINK\s+WRAP)\b",
        name, maxsplit=1, flags=re.I,
    )[0]
    name = re.sub(r"[-/]?\s*SERIES\s*\d+(?:\s*[-/]\s*\d+)?", " ", name, flags=re.I)
    name = re.sub(r"\bMIX(?:ED)?\s+BRAND\s*[- ]\s*TOYS\s*[- ]\s*MIXED\b", "", name, flags=re.I)

    for keyword, label in PRODUCT_NAME_ALIASES:
        name = re.sub(re.escape(keyword), label, name, flags=re.I)

    name = re.sub(r"\s*[-/]\s*", " ", name)
    name = re.sub(r"\bMIXED\b", "", name, flags=re.I)
    name = re.sub(r"\s+", " ", name).strip(" -/,，")

    # Some packing lists repeat the same product phrase on both sides of a separator.
    words = name.split()
    if len(words) % 2 == 0 and words[:len(words) // 2] == words[len(words) // 2:]:
        name = " ".join(words[:len(words) // 2])
    return name or source


def _pieces(explicit_pieces, quantity, specification):
    """The attachment carton count wins; calculation is a strict fallback only."""
    if explicit_pieces is not None:
        return explicit_pieces
    if quantity is None or not specification:
        return None
    calculated = quantity / specification
    return calculated if calculated.is_integer() else None


def _display_number(value) -> str:
    return str(int(value)) if float(value).is_integer() else str(value)


def _first(text: str, patterns: list[str]) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, re.I)
        if match:
            return match.group(1).strip()[:300]
    return ""


def _near_date(text: str, label: str) -> str:
    match = re.search(label, text, re.I)
    if not match:
        return ""
    tail = text[match.end():match.end() + 250]
    date = re.search(r"(\d{4}[/.-]\d{1,2}[/.-]\d{1,2}|\d{1,2}[/.-][A-Za-z]{3}[/.-]\d{2,4}|\d{1,2}/\d{1,2})(?:\s+\d{1,2}:\d{2})?", tail)
    return date.group().strip() if date else ""


def _normalize_container(value: str) -> str:
    match = re.search(r"(?:(\d+)\s*[xX*])?\s*(\d+)\s*(HQ|HC|GP)", value or "", re.I)
    return f"{match.group(1) or '1'}*{match.group(2)}{match.group(3).upper()}" if match else value


def _value_after_label(rows: list[list[str]], labels: tuple[str, ...]) -> str:
    for row in rows:
        for index, value in enumerate(row):
            if any(label.lower() in value.lower() for label in labels):
                for candidate in row[index + 1:]:
                    if candidate:
                        return candidate
    return ""


def _port(text: str) -> str:
    upper = text.upper()
    for token, value in (("YANTIAN", "盐田"), ("YICT", "盐田"), ("SHEKOU", "蛇口"), ("NANSHA", "南沙"), ("HUANGPU", "黄埔")):
        if token in upper:
            return value
    return ""


def _country(text: str) -> str:
    upper = text.upper()
    countries = (("UNITED STATES", "美国"), ("USA", "美国"), ("CANADA", "加拿大"), ("BRAZIL", "巴西"), ("UNITED KINGDOM", "英国"), ("AUSTRALIA", "澳大利亚"), ("GERMANY", "德国"), ("MEXICO", "墨西哥"))
    return next((value for token, value in countries if token in upper), "")


def _clean_address(value: str) -> str:
    if not value:
        return ""
    invalid = ("MOVEMENT TERM", "EXPORT LICENSE", "SOLID WOOD", "DISCHARGE ETA", "CARGO DESCRIPTION")
    upper = value.upper()
    if any(token in upper for token in invalid):
        return ""
    return value.strip(" ,;:")


def _row_label_value(rows: list[list], labels: tuple[str, ...]) -> str:
    for row in rows:
        for index, value in enumerate(row):
            if any(label.lower() in _text(value).lower() for label in labels):
                return next((_text(candidate) for candidate in row[index + 1:] if _text(candidate)), "")
    return ""


def _normalize_loading_port(value: str) -> str:
    upper = value.upper()
    if "YANTIAN" in upper or "YICT" in upper:
        return "盐田"
    if "SHEKOU" in upper:
        return "蛇口"
    if "NANSHA" in upper:
        return "南沙"
    return value
