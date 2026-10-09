from io import BytesIO
from openpyxl import load_workbook


def parse_factory_workbook(content: bytes) -> list[dict]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = [str(value or "").strip() for value in next(iterator, [])]
        required = ["英文名称", "中文简称", "本厂"]
        if any(column not in header for column in required):
            raise ValueError("表头必须包含：英文名称、中文简称、本厂")
        columns = [header.index(column) for column in required]
        rows = []
        for number, row in enumerate(iterator, 2):
            if not any(value not in (None, "") for value in row):
                continue
            values = [str(row[index] if index < len(row) and row[index] is not None else "").strip() for index in columns]
            if values[2] not in ("是", "否", "True", "False", "true", "false", "1", "0"):
                raise ValueError(f"第{number}行：本厂请填写是或否")
            rows.append({"englishName": values[0], "chineseShortName": values[1], "isLocal": values[2] in ("是", "True", "true", "1")})
            if len(rows) > 10000:
                raise ValueError("一次最多导入10000条映射")
        return rows
    finally:
        workbook.close()
