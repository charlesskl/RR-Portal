from copy import deepcopy
from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def _text(value):
    return str(value or "").strip()


def _weight(value, label):
    if value in (None, ""):
        return None
    try:
        result = float(value)
        if result < 0 or not __import__("math").isfinite(result):
            raise ValueError()
        return result
    except (TypeError, ValueError):
        raise ValueError(f"{label}须为非负数")


def normalized_packing_items(items):
    """Keep per-product quantity/net weight; common carton measurements count once."""
    result = deepcopy(items)
    groups = {}
    for index, item in enumerate(result):
        group = _text(item.get("packing_group"))
        if group:
            groups.setdefault(group, []).append(index)
    for group, indices in groups.items():
        if indices != list(range(indices[0], indices[-1] + 1)):
            raise ValueError(f"混装组 {group} 的货物须排列在一起")
        for key in ("pieces", "gross_weight", "volume"):
            values = [float(result[i][key]) for i in indices if result[i].get(key) not in (None, "")]
            if len(set(values)) > 1:
                raise ValueError(f"混装组 {group} 的共用件数、毛重或体积不一致，请只在第一行填写")
            result[indices[0]][key] = values[0] if values else None
            for i in indices[1:]:
                result[i][key] = None
    return result, groups


def build_huadeng_workbook(task, base_builder):
    from .shipment_export import _display_date, _display_deadline
    source = deepcopy(task)
    source["company"] = "Xingxin"
    source["specialRequirements"] = ""
    source["items"], groups = normalized_packing_items(list(task.get("items") or []))
    content, _ = base_builder(source)
    wb = load_workbook(BytesIO(content))
    ws = wb.active
    details = task.get("exportDetails") or {}
    mode = details.get("shipmentMode") or ("Warehouse" if task.get("warehouseGroups") else "Container")
    if mode not in {"Warehouse", "Container"}:
        raise ValueError("华登出货方式无效")
    warehouse = mode == "Warehouse"
    sequence = _text(details.get("sequenceNumber"))
    reference = f"{sequence or '待确认'}{'车' if warehouse else '柜'}"
    ws["A1"] = f"TO：{_text(details.get('recipient')) or '余振军'}"
    ws["F1"] = _text(details.get("pickupNote"))
    ws["F2"] = f"{_text(task.get('port'))} {_text(task.get('customer'))} {reference}".strip()
    ship_date = _display_date(task.get("plannedShipDate"))
    ws["E3"] = (f"{ship_date} 入{_text(details.get('warehouse')) or '待确认仓库'}  {_text(details.get('vehicleType'))}" if warehouse else
                 f"{ship_date} 出  {_text(task.get('containerType'))}  SO #{_text(task.get('soNumber'))}")
    ws["C5"] = "洋行"
    ws["F5"] = "货名"
    ws["N5"] = "客人PO#"
    for index, item in enumerate(source["items"], 6):
        ws.cell(index, 1, "华登")
    for indices in groups.values():
        if len(indices) < 2:
            continue
        first, last = indices[0] + 6, indices[-1] + 6
        for col in (10, 11, 13):
            ws.merge_cells(start_row=first, end_row=last, start_column=col, end_column=col)
        for col in (4, 5, 14, 15, 16):
            values = [_text(ws.cell(i + 6, col).value) for i in indices]
            if values[0] and all(not value or value == values[0] for value in values):
                ws.merge_cells(start_row=first, end_row=last, start_column=col, end_column=col)
    last_item = max(6, len(source["items"]) + 5)
    brand_row, total = last_item + 1, last_item + 2
    if details.get("brand"):
        ws.cell(brand_row, 6, "品牌：" + _text(details["brand"]))
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row > total:
            ws.unmerge_cells(str(merged))
    ws.delete_rows(total + 1, ws.max_row - total)
    pallet = _weight(details.get("palletWeight"), "卡板重量")
    for col in (9, 10, 11, 12, 13):
        letter = get_column_letter(col)
        extra = f"+{pallet}" if col == 11 and pallet is not None else ""
        ws.cell(total, col, f"=SUM({letter}6:{letter}{last_item}){extra}")
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def write(row, col, end, value):
        if end > col:
            ws.merge_cells(start_row=row, end_row=row, start_column=col, end_column=end)
        cell = ws.cell(row, col, value)
        cell.alignment = left
        cell.font = Font(name="宋体", size=12, bold=True)
        ws.row_dimensions[row].height = 30

    write(total, 1, 5, f"SI：{_display_deadline(task.get('siDeadline'))}")
    write(total + 1, 1, 5, f"截数期：{_display_deadline(task.get('cutoffDate'))}")
    write(total + 2, 1, 5, f"制表：{_text(details.get('preparedBy'))} {date.today().isoformat()}")
    if warehouse:
        write(total + 1, 6, 13, f"卡板说明：{_text(task.get('specialRequirements'))}")
        write(total + 2, 6, 13, f"收货人：{_text(details.get('consignee'))}")
        write(total + 2, 14, 17, f"收货人代码：{_text(details.get('consigneeCode'))}")
        end_row = total + 2
    else:
        filler = _weight(details.get("fillerWeight"), "填充物重量")
        container = _weight(details.get("containerWeight"), "柜重")
        write(total + 1, 8, 9, "填充物重量")
        ws.cell(total + 1, 10, filler)
        write(total + 2, 8, 9, "柜重")
        ws.cell(total + 2, 10, container)
        write(total + 3, 8, 9, "整个集装箱重量")
        if container is not None:
            ws.cell(total + 3, 10, f"=K{total}+J{total+1}+J{total+2}")
        write(total + 1, 14, 17, f"柜号：{_text(details.get('containerNumber')) or _text(task.get('transportReference'))}")
        write(total + 2, 14, 17, f"船封：{_text(details.get('sealNumber'))}")
        write(total + 4, 1, 4, f"贸易方式：{_text(details.get('tradeMode')) or '进料对口'}")
        write(total + 4, 5, 17, "发货人：东莞华登塑胶制品有限公司（4419943670）")
        write(total + 5, 1, 4, "拼箱：否")
        write(total + 5, 5, 13, f"收货人：{_text(details.get('consignee'))}")
        write(total + 5, 14, 17, f"收货人代码：{_text(details.get('consigneeCode'))}")
        write(total + 6, 1, 17, _text(task.get("specialRequirements")))
        end_row = total + 6
    ws.print_area = f"A1:{get_column_letter(ws.max_column)}{end_row}"
    out = BytesIO()
    wb.save(out)
    suffix = "交仓走货表" if warehouse else "走柜表"
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in reference)
    return out.getvalue(), f"华登_{_text(task.get('customer'))}_{safe}_{suffix}.xlsx"
