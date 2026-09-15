from __future__ import annotations

import base64
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

from .spreadsheets import _column_map, parse_business_spreadsheet


SUPPORTED_EXTENSIONS = (".xls", ".xlsx", ".xlsm")
CHANGED_FILL = PatternFill("solid", fgColor="FFF2CC")


def _key(value) -> str:
    return "".join(character for character in str(value or "").upper() if character.isalnum())


def _number(value) -> float:
    if value in (None, ""):
        return 0
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"库存数量格式无效：{value}") from exc


def _available_quantity(row: dict) -> float:
    """Inbound quantity is the source amount; outbound quantity is cumulative usage."""
    if row.get("inbound_quantity") not in (None, ""):
        return max(0, _number(row.get("inbound_quantity")) - _number(row.get("outbound_quantity")))
    return max(0, _number(row.get("inventory_quantity")))


def _date_sort_key(value) -> tuple:
    raw = str(value or "").strip()
    for parser in (date.fromisoformat, lambda text: datetime.strptime(text, "%Y/%m/%d").date()):
        try:
            return (0, parser(raw).isoformat())
        except ValueError:
            continue
    return (1, raw)


def _existing_outbound_date(value, reference_year: int) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    for format_string in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(raw, format_string).date()
        except ValueError:
            continue
    match = re.fullmatch(r"(\d{1,2})月(\d{1,2})日", raw)
    if match:
        return date(reference_year, int(match.group(1)), int(match.group(2)))
    return None


def _header_mapping(worksheet) -> dict:
    best_mapping = {}
    best_row = 0
    for row_index, values in enumerate(worksheet.iter_rows(min_row=1, max_row=min(40, worksheet.max_row), values_only=True), 1):
        mapping = _column_map(list(values))
        if len(mapping) > len(best_mapping):
            best_mapping = mapping
            best_row = row_index
    if not best_row:
        raise ValueError(f"工作表“{worksheet.title}”未找到库存表头")
    required = ("outbound_date", "outbound_quantity")
    missing = [field for field in required if field not in best_mapping]
    if missing:
        raise ValueError(f"工作表“{worksheet.title}”缺少出库日期或出库数量列")
    return best_mapping


def _convert_xls_to_xlsx(source: Path, output_directory: Path) -> Path:
    executable = shutil.which("soffice")
    if not executable:
        raise ValueError("旧版 .xls 库存表写回需要本机安装 LibreOffice")
    result = subprocess.run(
        [executable, "--headless", "--convert-to", "xlsx", "--outdir", str(output_directory), str(source)],
        capture_output=True, text=True, timeout=120, check=False,
    )
    converted = output_directory / f"{source.stem}.xlsx"
    if result.returncode or not converted.exists():
        raise ValueError(f"旧版库存表转换失败：{source.name}；{result.stderr or result.stdout}")
    _normalize_full_row_filter_refs(converted)
    return converted


def _normalize_full_row_filter_refs(path: Path) -> None:
    """LibreOffice may emit row-only filter refs that openpyxl cannot read."""
    normalized = path.with_name(f"normalized-{path.name}")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(normalized, "w") as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename.startswith("xl/worksheets/") and member.filename.endswith(".xml"):
                content = re.sub(
                    rb'ref="\$?(\d+):\$?(\d+)"',
                    lambda match: b'ref="A' + match.group(1) + b':XFD' + match.group(2) + b'"',
                    content,
                )
            target.writestr(member, content)
    normalized.replace(path)


def _convert_xlsx_to_xls(source: Path, output_directory: Path, original_name: str) -> Path:
    executable = shutil.which("soffice")
    if not executable:
        raise ValueError("旧版 .xls 库存表写回需要本机安装 LibreOffice")
    output_directory.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [executable, "--headless", "--convert-to", "xls:MS Excel 97", "--outdir", str(output_directory), str(source)],
        capture_output=True, text=True, timeout=120, check=False,
    )
    converted = output_directory / f"{source.stem}.xls"
    if result.returncode or not converted.exists():
        raise ValueError(f"旧版库存表还原失败：{original_name}；{result.stderr or result.stdout}")
    return converted


def _editable_path(source: Path, conversion_directory: Path) -> tuple[Path, bool]:
    if source.suffix.lower() == ".xls":
        return _convert_xls_to_xlsx(source, conversion_directory), True
    return source, False


def build_inventory_writeback(paths: list[Path], payload: dict) -> dict:
    task_id = str(payload.get("task_id") or "").strip()
    outbound_date = date.fromisoformat(str(payload.get("outbound_date") or date.today().isoformat()))
    task_items = list(payload.get("items") or [])
    if not task_id:
        raise ValueError("缺少走柜任务编号")
    if not task_items:
        raise ValueError("走柜任务没有货物明细")

    inventory_rows = []
    path_by_name = {path.name: path for path in paths if path.suffix.lower() in SUPPORTED_EXTENSIONS}
    for path in path_by_name.values():
        parsed = parse_business_spreadsheet(path, path.name)
        for row in parsed.get("rows", []):
            inventory_rows.append({**row, "source_file": path.name})

    remaining_by_row = {}
    for row in inventory_rows:
        row_key = (row["source_file"], row["source_sheet"], int(row["source_row"]))
        remaining_by_row[row_key] = _available_quantity(row)

    allocations = []
    for item_index, item in enumerate(task_items):
        item_outbound_date = date.fromisoformat(str(item.get("outbound_date") or outbound_date.isoformat()))
        requested = _number(item.get("quantity"))
        if requested <= 0:
            raise ValueError(f"第 {item_index + 1} 条明细的出库数量必须大于0")
        contract = _key(item.get("contract_number") or item.get("contractNumber"))
        product = _key(item.get("product_code") or item.get("productCode"))
        if not contract or not product:
            raise ValueError(f"第 {item_index + 1} 条明细缺少合同号或货号")
        candidates = [row for row in inventory_rows
                      if _key(row.get("contract_number")) == contract and _key(row.get("product_code")) == product]
        candidates.sort(key=lambda row: (_date_sort_key(row.get("inbound_date")), row["source_file"],
                                         row["source_sheet"], int(row["source_row"])))
        available = sum(remaining_by_row[(row["source_file"], row["source_sheet"], int(row["source_row"]))]
                        for row in candidates)
        if available < requested:
            raise ValueError(f"合同号 {item.get('contract_number')}、货号 {item.get('product_code')} 库存不足：需要 {requested:g}，可用 {available:g}")
        unallocated = requested
        for row in candidates:
            row_key = (row["source_file"], row["source_sheet"], int(row["source_row"]))
            available_in_row = remaining_by_row[row_key]
            if available_in_row <= 0:
                continue
            quantity = min(unallocated, available_in_row)
            remaining_by_row[row_key] -= quantity
            allocations.append({
                "item_index": item_index, "quantity": quantity, "file": row["source_file"],
                "sheet": row["source_sheet"], "row": int(row["source_row"]),
                "receipt_number": row.get("inbound_receipt_number") or "",
                "writeback_task_id": str(item.get("writeback_task_id") or task_id),
                "outbound_date": item_outbound_date.isoformat(),
                "outbound_trip": str(item.get("outbound_trip") or "").strip(),
                "remaining_after": remaining_by_row[row_key],
            })
            unallocated -= quantity
            if unallocated <= 0:
                break

    updates = {}
    for allocation in allocations:
        row_key = (allocation["file"], allocation["sheet"], allocation["row"])
        update = updates.setdefault(row_key, {"quantity": 0.0, "outbound_date": allocation["outbound_date"], "outbound_trips": []})
        update["quantity"] += allocation["quantity"]
        update["outbound_date"] = max(update["outbound_date"], allocation["outbound_date"])
        if allocation["outbound_trip"] and allocation["outbound_trip"] not in update["outbound_trips"]:
            update["outbound_trips"].append(allocation["outbound_trip"])

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    output_files = []
    with tempfile.TemporaryDirectory(prefix="voyageplex-writeback-convert-") as conversion_root:
        conversion_directory = Path(conversion_root)
        for filename in sorted({key[0] for key in updates}):
            source = path_by_name[filename]
            original_bytes = source.read_bytes()
            editable, was_xls = _editable_path(source, conversion_directory)
            workbook = openpyxl.load_workbook(editable, keep_vba=editable.suffix.lower() == ".xlsm")
            for (update_file, sheet_name, row_number), update in updates.items():
                if update_file != filename:
                    continue
                quantity = update["quantity"]
                worksheet = workbook[sheet_name]
                mapping = _header_mapping(worksheet)
                outbound_cell = worksheet.cell(row_number, mapping["outbound_quantity"] + 1)
                date_cell = worksheet.cell(row_number, mapping["outbound_date"] + 1)
                trip_cell = worksheet.cell(row_number, mapping["outbound_trip"] + 1) if "outbound_trip" in mapping else None
                inventory_cell = worksheet.cell(row_number, mapping["inventory_quantity"] + 1) if "inventory_quantity" in mapping else None
                outbound_cell.value = _number(outbound_cell.value) + quantity
                update_date = date.fromisoformat(update["outbound_date"])
                existing_date = _existing_outbound_date(date_cell.value, update_date.year)
                date_cell.value = max(update_date, existing_date) if existing_date else update_date
                outbound_cell.fill = CHANGED_FILL
                date_cell.fill = CHANGED_FILL
                if trip_cell is not None and update["outbound_trips"]:
                    existing_trips = [value.strip() for value in re.split(r"[、,，;；]", str(trip_cell.value or "")) if value.strip()]
                    trip_cell.value = "、".join(dict.fromkeys([*existing_trips, *update["outbound_trips"]]))
                    trip_cell.fill = CHANGED_FILL
                if inventory_cell is not None and not (isinstance(inventory_cell.value, str) and inventory_cell.value.startswith("=")):
                    if "inbound_quantity" in mapping:
                        inbound_cell = worksheet.cell(row_number, mapping["inbound_quantity"] + 1)
                        inventory_cell.value = max(0, _number(inbound_cell.value) - _number(outbound_cell.value))
                    else:
                        inventory_cell.value = max(0, _number(inventory_cell.value) - quantity)
                    inventory_cell.fill = CHANGED_FILL
            edited_xlsx = conversion_directory / f"edited-{source.stem}.xlsx"
            workbook.save(edited_xlsx)
            output_path = _convert_xlsx_to_xls(edited_xlsx, conversion_directory / f"xls-{source.stem}", filename) if was_xls else edited_xlsx
            output_files.append({
                "filename": filename,
                "content_base64": base64.b64encode(output_path.read_bytes()).decode("ascii"),
                "backup_filename": f"{source.stem}.回写前备份-{timestamp}{source.suffix}",
                "backup_base64": base64.b64encode(original_bytes).decode("ascii"),
            })
    return {"task_id": task_id, "outbound_date": outbound_date.isoformat(), "files": output_files,
            "allocations": allocations, "updated_rows": len(updates)}
