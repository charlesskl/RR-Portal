import tempfile
import uuid
import re
import os
import json
from datetime import datetime
from urllib.parse import quote
from pathlib import Path

from fastapi import FastAPI, File, Form, UploadFile
from fastapi.responses import Response

from .eml import parse_eml
from .rules import classify_email, filter_items_for_email, normalize_deadline, parse_body
from .attachments import parse_attachment
from .spreadsheets import parse_business_spreadsheet
from .shipment_export import build_completed_shipment_summary_workbook, build_inventory_adjustment_workbook, build_shipment_workbook
from .inventory_writeback import build_inventory_writeback
from .destination_country import infer_destination_country

PARSER_VERSION = "voyageplex-rules-0.1.2+destination-country"
app = FastAPI(title="VoyagePlex Email Parser", version=PARSER_VERSION)


def build_warehouse_groups(parsed_attachments: list[dict]) -> list[dict]:
    """Pair SO/PDF and PKL/Excel by filename number, then aggregate by warehouse."""
    by_reference: dict[str, dict] = {}
    for attachment in parsed_attachments:
        fields = attachment.get("fields", {})
        attachment_items = attachment.get("items", [])
        if not fields.get("delivery_address") and not attachment_items:
            continue
        filename = attachment["filename"]
        match = re.search(r"\b(\d{7,})\b", filename)
        reference = match.group(1) if match else Path(filename).stem
        entry = by_reference.setdefault(reference, {"reference": reference, "warehouse": "", "items": [], "source_files": []})
        entry["source_files"].append(filename)
        if fields.get("delivery_address"):
            entry["warehouse"] = fields["delivery_address"]
        entry["items"].extend(attachment_items)

    grouped: dict[str, dict] = {}
    for entry in by_reference.values():
        warehouse = entry["warehouse"] or "待确认仓库"
        group = grouped.setdefault(warehouse, {
            "warehouse": warehouse, "references": [], "items": [], "source_files": []
        })
        group["references"].append(entry["reference"])
        group["items"].extend(entry["items"])
        group["source_files"].extend(entry["source_files"])
    return list(grouped.values())


@app.get("/health")
def health():
    return {"status": "ok", "parser_version": PARSER_VERSION}


def _inventory_paths(root: Path) -> list[Path]:
    if not root.is_dir():
        return []
    return [path for path in sorted(root.iterdir())
            if not path.name.startswith("~$") and "回写前备份-" not in path.name
            and not path.name.startswith("库存扣减")
            and path.suffix.lower() in (".xls", ".xlsx", ".xlsm", ".csv")]


def _scan_inventory_paths(paths: list[Path], folder: str) -> dict:
    files = []
    for path in paths:
        stat = path.stat()
        item = {"filename": path.name, "modified_at": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
                "size": stat.st_size, "status": "ok", "kind": "", "sheets": [], "row_count": 0, "warnings": []}
        try:
            parsed = parse_business_spreadsheet(path, path.name)
            item.update({"kind": parsed.get("kind", "unknown"), "sheets": parsed.get("sheets", []),
                         "row_count": len(parsed.get("rows", [])), "warnings": parsed.get("warnings", [])})
        except Exception as exc:
            item.update({"status": "failed", "error": str(exc)})
        files.append(item)
    return {"folder": folder, "exists": True, "scanned_at": datetime.now().isoformat(timespec="seconds"),
            "files": files, "total_rows": sum(item["row_count"] for item in files),
            "successful_files": sum(item["status"] == "ok" for item in files)}


@app.get("/v1/local-inventory-scan")
def scan_local_inventory():
    root = Path(os.environ.get("VOYAGEPLEX_LOCAL_INVENTORY_DIR", "/Users/josie/Desktop/库存读取写回测试"))
    if not root.is_dir():
        return {"folder": str(root), "exists": False, "files": [], "total_rows": 0,
                "error": "本地库存文件夹不存在"}
    return _scan_inventory_paths(_inventory_paths(root), str(root))


@app.post("/v1/local-inventory-files/scan")
async def scan_uploaded_local_inventory(files: list[UploadFile] = File(...), folder: str = Form("本地库存文件夹")):
    with tempfile.TemporaryDirectory(prefix="voyageplex-local-inventory-") as temp_root:
        root = Path(temp_root)
        paths = []
        for file in files:
            filename = Path(file.filename or "inventory.xlsx").name
            if filename.startswith("~$") or "回写前备份-" in filename or filename.startswith("库存扣减") or Path(filename).suffix.lower() not in (".xls", ".xlsx", ".xlsm", ".csv"):
                continue
            path = root / filename
            path.write_bytes(await file.read())
            paths.append(path)
        return _scan_inventory_paths(paths, folder)


def _match_key(value) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def _match_inventory_paths(paths: list[Path], payload: dict) -> dict:
    inventory_rows = []
    for path in paths:
        parsed = parse_business_spreadsheet(path, path.name)
        for row in parsed.get("rows", []):
            inventory_rows.append({**row, "source_file": path.name})
    customer = _match_key(payload.get("customer"))
    results = []
    for index, item in enumerate(payload.get("items") or []):
        contract = _match_key(item.get("contract_number") or item.get("contractNumber"))
        product = _match_key(item.get("product_code") or item.get("productCode"))
        matches = [row for row in inventory_rows
                   if contract and product and _match_key(row.get("contract_number")) == contract
                   and _match_key(row.get("product_code")) == product
                   and (not customer or customer in _match_key(row.get("source_sheet")))]
        if not matches:
            matches = [row for row in inventory_rows if contract and product
                       and _match_key(row.get("contract_number")) == contract
                       and _match_key(row.get("product_code")) == product]
        locations = sorted({str(row.get("storage_location") or "").strip() for row in matches if row.get("storage_location")})
        receipts = sorted({str(row.get("inbound_receipt_number") or "").strip() for row in matches if row.get("inbound_receipt_number")})
        countries = sorted({str(row.get("country") or "").strip() for row in matches if row.get("country")})
        customers = sorted({str(row.get("customer") or "").strip() for row in matches if row.get("customer")})
        sources = sorted({f'{row.get("source_file")} / {row.get("source_sheet")} / 第{row.get("source_row")}行' for row in matches})
        results.append({"index": index, "match_count": len(matches), "location": "、".join(locations),
                        "receipt_number": "、".join(receipts), "country": "、".join(countries),
                        "customer_name": "、".join(customers), "sources": sources,
                        "status": "已唯一匹配" if len(matches) == 1 else ("多条匹配，需确认" if matches else "未匹配")})
    return {"items": results, "scanned_files": len({row["source_file"] for row in inventory_rows}),
            "inventory_rows": len(inventory_rows)}


@app.post("/v1/local-inventory-match")
def match_local_inventory(payload: dict):
    root = Path(os.environ.get("VOYAGEPLEX_LOCAL_INVENTORY_DIR", "/Users/josie/Desktop/库存读取写回测试"))
    if not root.is_dir():
        return {"error": "本地库存文件夹不存在", "items": []}
    return _match_inventory_paths(_inventory_paths(root), payload)


@app.post("/v1/local-inventory-files/match")
async def match_uploaded_local_inventory(payload: str = Form(...), files: list[UploadFile] = File(...)):
    request_payload = json.loads(payload)
    with tempfile.TemporaryDirectory(prefix="voyageplex-local-inventory-") as temp_root:
        root = Path(temp_root)
        paths = []
        for file in files:
            filename = Path(file.filename or "inventory.xlsx").name
            if filename.startswith("~$") or "回写前备份-" in filename or filename.startswith("库存扣减") or Path(filename).suffix.lower() not in (".xls", ".xlsx", ".xlsm", ".csv"):
                continue
            path = root / filename
            path.write_bytes(await file.read())
            paths.append(path)
        return _match_inventory_paths(paths, request_payload)


@app.post("/v1/local-inventory-files/writeback")
async def writeback_uploaded_local_inventory(payload: str = Form(...), files: list[UploadFile] = File(...)):
    request_payload = json.loads(payload)
    with tempfile.TemporaryDirectory(prefix="voyageplex-local-inventory-") as temp_root:
        root = Path(temp_root)
        paths = []
        for file in files:
            filename = Path(file.filename or "inventory.xlsx").name
            if filename.startswith("~$") or "回写前备份-" in filename or filename.startswith("库存扣减") or Path(filename).suffix.lower() not in (".xls", ".xlsx", ".xlsm"):
                continue
            path = root / filename
            path.write_bytes(await file.read())
            paths.append(path)
        if not paths:
            return {"error": "所选文件夹中没有可写回的库存表"}
        try:
            return build_inventory_writeback(paths, request_payload)
        except ValueError as exc:
            return Response(content=json.dumps({"error": str(exc)}, ensure_ascii=False), media_type="application/json", status_code=400)


@app.post("/v1/shipment-export")
async def export_shipment(payload: dict):
    content, filename = build_shipment_workbook(payload)
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@app.post("/v1/inventory-adjustment-export")
async def export_inventory_adjustment(payload: dict):
    content, filename = build_inventory_adjustment_workbook(payload)
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@app.post("/v1/completed-shipment-summary-export")
async def export_completed_shipment_summary(payload: dict):
    content, filename = build_completed_shipment_summary_workbook(payload)
    return Response(content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


@app.post("/v1/email-batches/parse")
async def parse_email_batch(files: list[UploadFile] = File(...)):
    batch_id = str(uuid.uuid4())
    results = []
    with tempfile.TemporaryDirectory(prefix=f"voyageplex-{batch_id}-") as root:
        for index, upload in enumerate(files):
            item = {"index": index, "filename": upload.filename or f"email-{index}.eml"}
            try:
                if not item["filename"].lower().endswith(".eml"):
                    raise ValueError("只支持 .eml 邮件文件")
                raw = await upload.read()
                if not raw:
                    raise ValueError("文件为空")
                message = parse_eml(raw, Path(root) / f"item-{index}")
                fields = parse_body(f'{message["subject"]}\n{message["body_text"]}')
                fallback_dates = {
                    "si_deadline": fields.pop("_fallback_si_deadline", ""),
                    "cutoff_date": fields.pop("_fallback_cutoff_date", ""),
                }
                date_warnings = fields.pop("_date_warnings", [])
                shipment_type = classify_email(message["subject"], message["body_text"])
                attachment_results = []
                attachment_details = []
                warnings = list(date_warnings)
                items = []
                attachment_date_candidates = {"si_deadline": [], "cutoff_date": []}
                supported_attachment_count = 0
                successful_attachment_count = 0
                for attachment in message["attachments"]:
                    suffix = Path(attachment["filename"]).suffix.lower()
                    is_supported = suffix in (".xlsx", ".xlsm", ".xls", ".pdf", ".docx")
                    if is_supported:
                        supported_attachment_count += 1
                    try:
                        parsed_attachment = parse_attachment(Path(attachment["stored_path"]), attachment["filename"])
                    except Exception as exc:
                        attachment_results.append({
                            "filename": attachment["filename"], "kind": "failed", "status": "failed",
                            "fields": {}, "item_count": 0, "error": str(exc),
                        })
                        warnings.append(f'已跳过无法解析附件：{attachment["filename"]}（{exc}）')
                        continue
                    if is_supported:
                        successful_attachment_count += 1
                    parsed_attachment["items"] = filter_items_for_email(
                        parsed_attachment["items"], shipment_type,
                        str(parsed_attachment["fields"].get("loading_factory", "")),
                    )
                    attachment_results.append({
                        "filename": attachment["filename"],
                        "kind": parsed_attachment["kind"],
                        "status": "parsed",
                        "fields": parsed_attachment["fields"],
                        "item_count": len(parsed_attachment["items"]),
                    })
                    attachment_details.append({
                        "filename": attachment["filename"],
                        "kind": parsed_attachment["kind"],
                        "fields": parsed_attachment["fields"],
                        "items": parsed_attachment["items"],
                    })
                    warnings.extend(parsed_attachment["warnings"])
                    items.extend({**detail, "source_file": attachment["filename"]} for detail in parsed_attachment["items"])
                    for key, value in parsed_attachment["fields"].items():
                        if key == "po_so_mapping" or value in (None, "", []):
                            continue
                        target_key = "cutoff_date" if key == "customs_cutoff" else key
                        if target_key in ("si_deadline", "cutoff_date"):
                            value = normalize_deadline(str(value))
                            if value and value not in attachment_date_candidates[target_key]:
                                attachment_date_candidates[target_key].append(value)
                            continue
                        # The latest instructions in the message body take precedence over
                        # attachments; attachments only fill fields missing from that section.
                        if value and not fields.get(target_key):
                            fields[target_key] = value
                for key, values in attachment_date_candidates.items():
                    used_attachment = bool(values) and not fields.get(key)
                    if used_attachment:
                        fields[key] = min(values)
                    if used_attachment and len(values) > 1:
                        label = "SI截止" if key == "si_deadline" else "截关"
                        warnings.append(
                            f"附件中{label}识别到多个时间（{'、'.join(values)}），"
                            f"已采用最早时间，请人工确认"
                        )
                for key, value in fallback_dates.items():
                    if value and not fields.get(key):
                        fields[key] = value
                if supported_attachment_count > 0 and successful_attachment_count == 0:
                    raise ValueError("邮件中的业务附件均无法解析")
                warehouse_groups = build_warehouse_groups(attachment_details)
                fields["destination_country"] = infer_destination_country(
                    f'{message["subject"]}\n{message["body_text"]}', attachment_details)
                item.update({
                    "status": "parsed",
                    "fingerprint": message["fingerprint"],
                    "message": {k: message[k] for k in ("message_id", "subject", "sender", "received_at")},
                    "attachments": [{k: v for k, v in attachment.items() if k != "stored_path"} for attachment in message["attachments"]],
                    "attachment_results": attachment_results,
                    "fields": fields,
                    "shipment_type": shipment_type,
                    "items": items,
                    "warehouse_groups": warehouse_groups,
                    "so_numbers": [reference for group in warehouse_groups for reference in group["references"]],
                    "warnings": warnings,
                })
            except Exception as exc:
                item.update({"status": "failed", "error": str(exc)})
            results.append(item)
    return {
        "batch_id": batch_id,
        "parser_version": PARSER_VERSION,
        "total": len(results),
        "parsed": sum(x["status"] == "parsed" for x in results),
        "failed": sum(x["status"] == "failed" for x in results),
        "items": results,
    }


@app.post("/v1/spreadsheet-batches/parse")
async def parse_spreadsheet_batch(files: list[UploadFile] = File(...)):
    results = []
    with tempfile.TemporaryDirectory(prefix="voyageplex-sheets-") as root:
        for index, upload in enumerate(files):
            item = {"index": index, "filename": upload.filename or f"sheet-{index}.xlsx"}
            try:
                suffix = Path(item["filename"]).suffix.lower()
                if suffix not in (".xlsx", ".xlsm", ".xls", ".csv"):
                    raise ValueError("仅支持 xlsx、xls、csv 表格")
                data = await upload.read()
                path = Path(root) / f"{index}{suffix}"
                path.write_bytes(data)
                parsed = parse_business_spreadsheet(path, item["filename"])
                item.update({"status": "parsed", **parsed})
            except Exception as exc:
                item.update({"status": "failed", "error": str(exc), "rows": [], "warnings": []})
            results.append(item)
    return {"parser_version": PARSER_VERSION, "total": len(results),
            "parsed": sum(item["status"] == "parsed" for item in results),
            "failed": sum(item["status"] == "failed" for item in results), "items": results}


def _mapping_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _mapping_row(group, source, customer, code, name="", owner="", place="", note="", excluded=False):
    code = _mapping_text(code).replace("\n", "")
    reserved = {"货号", "产品货号", "序号", "合计", "半成品", "ZURU", "MOOSE", "CEPIA", "LIFELINES", "TIGERHEAD", "TOMY", "ZANAOON"}
    if not code or code.upper() in reserved:
        return None
    return {"group_name": group, "inspection_source": source, "is_excluded": excluded,
            "customer": _mapping_text(customer), "product_code": code,
            "product_name": _mapping_text(name), "owner": _mapping_text(owner),
            "production_place": _mapping_text(place), "note": _mapping_text(note)}


def parse_inspection_mapping(path: Path) -> list[dict]:
    import openpyxl
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    result = []
    for sheet in workbook.worksheets:
        title = sheet.title.strip()
        rows = list(sheet.iter_rows(max_col=min(sheet.max_column, 20), values_only=True))
        if title.startswith("兴信B车间"):
            for start in (0, 4, 8, 12, 16):
                customer = _mapping_text(rows[0][start]) if rows and start < len(rows[0]) else ""
                for values in rows[2:]:
                    row = _mapping_row("兴信B车间", "兴信验货总结表", customer,
                                       values[start] if start < len(values) else None,
                                       values[start + 1] if start + 1 < len(values) else None,
                                       "龙丽娟", "兴信B车间")
                    if row: result.append(row)
        elif title.startswith("兴信A车间石玉珍"):
            customer = ""
            for values in rows[1:]:
                customer = _mapping_text(values[0]) or customer
                raw = _mapping_text(values[1])
                code, _, name = raw.partition("/")
                row = _mapping_row("兴信A车间 · 石玉珍", "兴信验货总结表", customer, code, name,
                                   "石玉珍", "兴信A车间")
                if row: result.append(row)
        elif title.startswith("兴信A车间李诗妍"):
            customer = ""
            for values in rows[2:]:
                customer = _mapping_text(values[1]) or customer
                place = _mapping_text(values[3]) or "兴信A车间"
                row = _mapping_row("兴信A车间 · 李诗妍", "兴信验货总结表", customer, values[2], "",
                                   "李诗妍", place)
                if row: result.append(row)
        elif "华嘉负责货号" in title:
            customer = ""
            for values in rows[2:]:
                customer = _mapping_text(values[1]) or customer
                row = _mapping_row("华嘉", "兴信验货总结表", customer, values[2], "", "华嘉",
                                   values[3], "兴信表")
                if row: result.append(row)
        elif title.startswith("湖南钟雅娴"):
            customer = ""
            place = ""
            for values in rows[1:]:
                place = _mapping_text(values[1]) or place or "湖南"
                customer = _mapping_text(values[2]) or customer
                owner = _mapping_text(values[7]) or _mapping_text(values[6]) or "钟雅娴"
                row = _mapping_row("湖南 · 钟雅娴", "湖南验货总结表", customer, values[3], values[4], owner, place)
                if row: result.append(row)
        elif title.startswith("湖南砚秋") or title.startswith("湖南Evol"):
            owner = "吕宇祥" if "吕宇祥" in title else "杨海彬"
            group = f"湖南 · {owner}"
            for start in (0, 4, 8, 12, 16):
                if not rows or start >= len(rows[0]): continue
                customer = _mapping_text(rows[0][start])
                for values in rows[2:]:
                    row = _mapping_row(group, "湖南验货总结表", customer,
                                       values[start] if start < len(values) else None,
                                       values[start + 1] if start + 1 < len(values) else None,
                                       owner, "湖南")
                    if row: result.append(row)
        elif title.startswith("河源"):
            for values in rows:
                note = "自行做柜"
                if len(values) > 6 and _mapping_text(values[6]): note += "；" + _mapping_text(values[6])
                row = _mapping_row("河源", "不参与验货检查", "ZURU", values[0], values[1], values[2],
                                   "河源", note, True)
                if row: result.append(row)
        elif title.startswith("华登"):
            owners = {"东莞华登": "王远露", "湖南华登": _mapping_text(rows[2][5]) if len(rows) > 2 else ""}
            customer = _mapping_text(rows[0][0]) if rows else "ZURU"
            for values in rows[2:]:
                place = _mapping_text(values[2])
                row = _mapping_row("华登", "华登验货总结表", customer, values[0], values[1],
                                   owners.get(place, ""), place)
                if row: result.append(row)
    unique = {}
    for row in result:
        unique[(row["group_name"], row["product_code"])] = row
    return list(unique.values())


@app.post("/v1/inspection-mappings/parse")
async def parse_inspection_mapping_file(file: UploadFile = File(...)):
    filename = file.filename or "跟单负责货号.xlsx"
    if Path(filename).suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("验货映射目前只支持 xlsx、xlsm 文件")
    with tempfile.TemporaryDirectory(prefix="voyageplex-mapping-") as root:
        path = Path(root) / Path(filename).name
        path.write_bytes(await file.read())
        rows = parse_inspection_mapping(path)
    if not rows:
        raise ValueError("未从跟单负责货号表中识别到货号")
    return {"filename": filename, "total": len(rows), "rows": rows}
