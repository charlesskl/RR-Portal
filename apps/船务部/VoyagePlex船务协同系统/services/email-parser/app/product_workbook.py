"""Read product records from VoyagePlex shipment workbooks."""

from io import BytesIO
import re

from openpyxl import load_workbook


HEADERS = ("序号", "客户", "货号", "货名 / 装箱规格", "玩具类别", "每箱毛重", "每箱净重")
SPEC_PATTERN = re.compile(r"(?:^|\s)(\d+)\s*个\s*/\s*箱\s*$")


def parse_product_workbook(content: bytes, filename: str) -> list[dict]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError(f"{filename} 不是有效的走柜表 Excel 文件") from error
    try:
        if "柜单" not in workbook.sheetnames:
            raise ValueError(f"{filename} 缺少“柜单”工作表")
        sheet = workbook["柜单"]
        header = [str(sheet.cell(5, index).value or "").strip() for index in range(1, 18)]
        if any(value not in header for value in HEADERS):
            raise ValueError(f"{filename} 的柜单表头与系统走柜表不符")
        columns = {value: header.index(value) for value in HEADERS}
        rows = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=6, values_only=True), start=6):
            sequence = cells[columns["序号"]] if len(cells) > columns["序号"] else None
            if sequence is None or not str(sequence).strip().isdigit():
                break

            def cell(label: str) -> str:
                index = columns[label]
                return str(cells[index]).strip() if index < len(cells) and cells[index] is not None else ""

            combined_name = cell("货名 / 装箱规格")
            match = SPEC_PATTERN.search(combined_name)
            quantity = int(match.group(1)) if match else None
            name = combined_name[:match.start()].strip() if match else combined_name

            def weight(label: str) -> float | None:
                raw = cell(label)
                if not raw:
                    return None
                try:
                    value = float(raw)
                    return value if value > 0 else None
                except ValueError:
                    return None

            warnings = []
            if not cell("货号"):
                warnings.append("货号为空")
            if quantity is None:
                warnings.append("未识别装箱规格，请核对")
            rows.append({
                "filename": filename, "rowNumber": row_number,
                "customer": cell("客户"), "productCode": cell("货号"),
                "productName": name, "quantityPerBox": quantity,
                "toyCategory": cell("玩具类别"),
                "grossWeightPerBox": weight("每箱毛重"),
                "netWeightPerBox": weight("每箱净重"), "warnings": warnings,
            })
        return rows
    finally:
        workbook.close()
