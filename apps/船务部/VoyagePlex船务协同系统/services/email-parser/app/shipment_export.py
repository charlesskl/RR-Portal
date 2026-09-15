from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


CONSIGNEES = {
    "ZURU": ("ZURU INC.", "66934704"),
    "TOMY": ("TOMY INTERNATIONAL, INC.", "51-0305290"),
}

# 柜单备注使用基础货号匹配：货号后的字母及后缀不参与厂区和品牌判断。
HUADENG_PRODUCT_CODES = {
    "92150", "92107", "92104", "15706", "15789", "15759", "9298",
    "15756", "15733", "92108", "15797", "92106", "157124", "157101",
    "9577", "15704", "77907",
}
HUAKANG_PRODUCT_CODES = {
    "15746", "15751", "15755", "15760", "15749", "15754", "77903",
}


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float | int | str:
    if value in (None, ""):
        return ""
    try:
        number = float(value)
        return int(number) if number.is_integer() else number
    except (TypeError, ValueError):
        return _text(value)


def _customer_key(task: dict) -> str:
    customer = _text(task.get("customer")).upper()
    return "TOMY" if "TOMY" in customer else "ZURU"


def _display_date(value: Any) -> str:
    raw = _text(value)
    if not raw:
        return ""
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        return f"{parsed.month}月{parsed.day}日"
    except ValueError:
        return raw


def _display_deadline(value: Any) -> str:
    raw = _text(value)
    if not raw:
        return ""
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        time = f" {parsed.hour:02d}:{parsed.minute:02d}" if "T" in raw or " " in raw else ""
        return f"{parsed.month}月{parsed.day}日{time}"
    except ValueError:
        return raw


def _item_value(item: dict, *keys: str) -> Any:
    for key in keys:
        value = item.get(key)
        if value not in (None, ""):
            return value
    return ""


def _base_product_code(value: Any) -> str:
    match = re.match(r"\s*(\d+)", _text(value))
    return match.group(1) if match else ""


def _factory_note(item: dict) -> str:
    product_code = _base_product_code(_item_value(item, "product_code", "productCode"))
    if product_code in HUADENG_PRODUCT_CODES:
        return "华登"
    if product_code in HUAKANG_PRODUCT_CODES:
        return "华康"
    return "兴信"


def _brand_note_lines(items: list[dict]) -> list[str]:
    matched: set[tuple[str, str]] = set()
    for item in items:
        product_code = _base_product_code(_item_value(item, "product_code", "productCode"))
        if product_code == "15756":
            matched.add(("15756品牌", "SPONGEBOB SQUAREPANTS"))
        elif product_code in {"92107", "92108"}:
            matched.add((f"{product_code}品牌", "ZURU/BABYCORNS"))
        elif product_code.startswith("157"):
            matched.add(("157开头的品牌", "ZURU/FUGGLER"))
        elif product_code.startswith("95"):
            matched.add(("95开头品牌", "ZURU/PetsAlive"))
        elif product_code.startswith("92"):
            matched.add(("92开头品牌", "ZURU/RAinBocoRns"))
        elif product_code.startswith("7"):
            matched.add(("7字开头的品牌", "ZURU"))

    order = {
        "157开头的品牌": 1, "95开头品牌": 2, "7字开头的品牌": 3,
        "92开头品牌": 4, "92107品牌": 5, "92108品牌": 6, "15756品牌": 7,
    }
    return [f"{label}：{brand}" for label, brand in sorted(matched, key=lambda value: order[value[0]])]


def _product_name_spec(item: dict) -> str:
    name = _text(item.get("product_name"))
    spec = _number(item.get("spec"))
    return f"{name}{' ' if name else ''}{spec}个/箱" if spec != "" else name


def _supplier_names(items: list[dict]) -> list[str]:
    names = []
    for item in items:
        supplier = _text(item.get("supplier"))
        if supplier and supplier not in names:
            names.append(supplier)
    return names


def _split_shipment_notes(notes: str, items: list[dict]) -> tuple[str, str]:
    """Split reviewed email remarks into the cabinet header and Xingxin-only footer."""
    lines = [line.strip(" ，,。;；") for line in notes.splitlines() if line.strip()]
    top_tokens = ("客上柜", "客上车", "拼柜", "拖柜", "深圳报关", "报关行")
    customer_pickup = any(token in notes for token in ("客上柜", "客上车"))
    top_lines = [line for line in lines if any(token in line for token in top_tokens)]
    ordinary_lines = [line for line in lines if line not in top_lines]

    suppliers = _supplier_names(items)
    xingxin_loads = not suppliers or any(re.search(r"兴信|新信|hanson", name, re.I) for name in suppliers)
    external = [name for name in suppliers if not re.search(r"兴信|新信|hanson", name, re.I)]

    if customer_pickup:
        customs_match = re.search(r"([\u4e00-\u9fffA-Za-z0-9]+报关行)", notes)
        customs = customs_match.group(1) if customs_match else "待确认报关行"
        top_lines = [f"{customs}，客上柜，拼柜，深圳报关"]
    elif not top_lines:
        if xingxin_loads and external:
            top_lines.append(f"兴信拖柜，{'、'.join(external)}送兴信拼柜，深圳报关")
        elif external and not xingxin_loads:
            top_lines.append(f"送{'、'.join(external)}拼柜，深圳报关")

    # 外厂做柜不带底部普通邮件备注；只有兴信本厂做柜（整柜或散货）才输出。
    return "\n".join(dict.fromkeys(top_lines)), "\n".join(ordinary_lines) if xingxin_loads else ""


def build_shipment_workbook(task: dict) -> tuple[bytes, str]:
    """Create one cabinet sheet per task, following the existing ZURU/TOMY layout."""
    customer_key = _customer_key(task)
    consignee, consignee_code = CONSIGNEES[customer_key]
    items = list(task.get("items") or [])
    notes = _text(task.get("specialRequirements"))
    top_note, bottom_note = _split_shipment_notes(notes, items)

    wb = Workbook()
    ws = wb.active
    ws.title = "柜单"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.2
    ws.page_margins.top = ws.page_margins.bottom = 0.35

    warehouse_export = _text(task.get("exportVariant")) == "warehouse"
    widths = [9, 6, 9, 16, 13, 27, 10, 11, 11, 9, 10, 10, 11, 16, 11, 11, 11]
    if warehouse_export:
        widths.append(18)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    yellow = PatternFill("solid", fgColor="FFF200")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.merge_cells("A1:D1")
    ws["A1"] = "TO：廖平/王军"
    ws["A1"].font = Font(size=14, bold=True)
    ws.merge_cells("F1:L1")
    ws["F1"] = top_note
    ws["F1"].fill = yellow
    ws["F1"].font = Font(size=12, bold=True)
    ws["F1"].alignment = center

    port = _text(task.get("port")) or "待确认港口"
    ws.merge_cells("F2:L2")
    ws["F2"] = f"{port} {customer_key} 走柜表"
    ws["F2"].font = Font(size=20, bold=True)
    ws["F2"].alignment = center

    ship_date = _display_date(task.get("plannedShipDate"))
    container_type = _text(task.get("containerType")) or "待确认柜型"
    so_number = (_text(task.get("soNumber")).lstrip("#").strip() or "待确认SO")
    ws.merge_cells("E3:M3")
    ws["E3"] = f"{ship_date} 出    {container_type}    SO #{so_number}"
    ws["E3"].font = Font(size=15, bold=True)
    ws["E3"].alignment = center

    headers = ["备注", "序号", "客户", "合同", "货号", "货名 / 装箱规格", "国家", "玩具类别", "数量", "件数", "毛重", "净重", "体积", "客户PO#", "每单总件数", "每箱毛重", "每箱净重"]
    if warehouse_export:
        headers.append("放货区")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(5, col, header)
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[5].height = 32

    first_item_row = 6
    for offset, item in enumerate(items):
        row = first_item_row + offset
        values = [
            _factory_note(item), offset + 1,
            customer_key, _item_value(item, "contract_number", "contractNumber"),
            _item_value(item, "product_code", "productCode"), _product_name_spec(item),
            item.get("country", ""), item.get("category", ""), _number(item.get("quantity")),
            _number(item.get("pieces")), _number(item.get("gross_weight")), _number(item.get("net_weight")), _number(_item_value(item, "volume", "cbm")),
            _item_value(item, "customer_po", "customerPo"), _number(_item_value(item, "order_total_pieces", "pieces")),
            _number(item.get("gross_weight_per_box")), _number(item.get("net_weight_per_box")),
        ]
        if warehouse_export:
            values.append(_item_value(item, "warehouse_location", "warehouseLocation"))
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.border = border
            cell.alignment = center if col not in (1, 6, 18) else left
        ws.row_dimensions[row].height = 34

    last_item_row = max(first_item_row, first_item_row + len(items) - 1)
    if not items:
        for col in range(1, len(headers) + 1):
            ws.cell(first_item_row, col).border = border
    brand_row = last_item_row + 1
    brand_notes = _brand_note_lines(items) if customer_key == "ZURU" else (["品牌：TOMY"] if items else [])
    ws.cell(brand_row, 6, "\n".join(brand_notes))
    ws.cell(brand_row, 6).font = Font(bold=True)
    ws.cell(brand_row, 6).alignment = left
    ws.row_dimensions[brand_row].height = max(34, 18 * len(brand_notes))
    for col in range(1, len(headers) + 1):
        ws.cell(brand_row, col).border = border

    total_row = brand_row + 1
    ws.merge_cells(start_row=total_row, start_column=7, end_row=total_row, end_column=8)
    ws.cell(total_row, 7, "合计")
    ws.cell(total_row, 7).font = Font(size=14, bold=True)
    ws.cell(total_row, 7).alignment = center
    if items:
        ws.cell(total_row, 9, f"=SUM(I{first_item_row}:I{last_item_row})")
        ws.cell(total_row, 10, f"=SUM(J{first_item_row}:J{last_item_row})")
        ws.cell(total_row, 13, f"=SUM(M{first_item_row}:M{last_item_row})")
    for col in range(7, 14):
        ws.cell(total_row, col).border = border
        ws.cell(total_row, col).alignment = center

    info_row = total_row + 1
    ws.merge_cells(start_row=info_row, start_column=1, end_row=info_row, end_column=5)
    ws.cell(info_row, 1, f"SI：{_display_deadline(task.get('siDeadline'))}")
    ws.merge_cells(start_row=info_row, start_column=9, end_row=info_row, end_column=11)
    ws.cell(info_row, 9, "填充物重量：")
    ws.merge_cells(start_row=info_row, start_column=14, end_row=info_row, end_column=len(headers))
    ws.cell(info_row, 14, "柜号：\n船封：")

    info_row2 = info_row + 1
    ws.merge_cells(start_row=info_row2, start_column=1, end_row=info_row2, end_column=5)
    ws.cell(info_row2, 1, f"截数期：{_display_deadline(task.get('cutoffDate'))}")
    ws.merge_cells(start_row=info_row2, start_column=9, end_row=info_row2, end_column=11)
    ws.cell(info_row2, 9, "柜重：")

    info_row3 = info_row2 + 1
    ws.merge_cells(start_row=info_row3, start_column=1, end_row=info_row3, end_column=5)
    ws.cell(info_row3, 1, f"制表：admin    {date.today().isoformat()}")
    ws.merge_cells(start_row=info_row3, start_column=9, end_row=info_row3, end_column=11)
    ws.cell(info_row3, 9, "整个集装箱重量：")

    footer_row = info_row3 + 2
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=len(headers))
    ws.cell(footer_row, 1, "贸易方式：进料对口      发货人：东莞兴信塑胶制品有限公司（4419946995）")
    ws.cell(footer_row, 1).font = Font(bold=True)
    footer_row2 = footer_row + 1
    ws.merge_cells(start_row=footer_row2, start_column=1, end_row=footer_row2, end_column=len(headers))
    ws.cell(footer_row2, 1, f"收货人：{consignee}      收货人代码：{consignee_code}")
    ws.cell(footer_row2, 1).font = Font(size=13, bold=True)

    if bottom_note:
        note_row = footer_row2 + 2
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 1, end_column=len(headers))
        ws.cell(note_row, 1, bottom_note)
        ws.cell(note_row, 1).fill = yellow
        ws.cell(note_row, 1).font = Font(bold=True, color="C00000")
        ws.cell(note_row, 1).alignment = left
        ws.row_dimensions[note_row].height = 28

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.alignment == Alignment():
                cell.alignment = left
    ws.print_area = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    safe_so = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in so_number)[:60]
    suffix = "仓务走柜表" if warehouse_export else "走柜表"
    filename = f"{customer_key}_{safe_so}_{suffix}.xlsx"
    output = BytesIO()
    wb.save(output)
    return output.getvalue(), filename


def build_inventory_adjustment_workbook(payload: dict) -> tuple[bytes, str]:
    """Create one concise write-back list for a completed-date range."""
    tasks = list(payload.get("tasks") or [payload])
    wb = Workbook()
    ws = wb.active
    ws.title = "库存扣减写回表"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    headers = ["洋行", "合同号", "客户名称", "走货国家", "货号", "入库单号", "放货区", "出库日期", "出库数量", "出柜车次"]
    widths = [14, 18, 28, 16, 22, 20, 26, 14, 14, 18]
    thin = Side(style="thin", color="B7BBC5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    changed_fill = PatternFill("solid", fgColor="FFF2CC")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center
        ws.column_dimensions[get_column_letter(col)].width = width
    row = 2
    for task in tasks:
        outbound_date = task.get("completedDate") or task.get("plannedShipDate") or ""
        transport = _item_value(task, "transportReference", "vehicleContainerNumber") or "待补充"
        for item in list(task.get("items") or []):
            values = [task.get("customer", ""), _item_value(item, "contract_number", "contractNumber"),
                      _item_value(item, "inventory_customer_name"), _item_value(item, "inventory_country", "country"),
                      _item_value(item, "product_code", "productCode"), _item_value(item, "inventory_receipt_number"),
                      _item_value(item, "warehouse_location", "warehouseLocation"), outbound_date,
                      _number(item.get("quantity")), transport]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row, col, value); cell.border = border
                cell.alignment = left if col in (2, 3, 5, 6, 7, 10) else center
                if col in (8, 9, 10): cell.fill = changed_fill
            ws.row_dimensions[row].height = 30
            row += 1
    ws.auto_filter.ref = f"A1:J{max(2, ws.max_row)}"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    date_range = f"{_text(payload.get('from'))}_{_text(payload.get('to'))}".strip("_") or date.today().isoformat()
    output = BytesIO()
    wb.save(output)
    return output.getvalue(), f"{date_range}_库存扣减写回表.xlsx"


def build_completed_shipment_summary_workbook(payload: dict) -> tuple[bytes, str]:
    """Create a concise task-level summary for a completed-date range."""
    tasks = list(payload.get("tasks") or [])
    wb = Workbook()
    ws = wb.active
    ws.title = "走柜任务汇总"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    headers = ["完成日期", "洋行", "SO号", "柜型", "走货日期", "装货港", "货物明细", "总数量", "总件数", "总体积(CBM)"]
    widths = [14, 14, 34, 14, 14, 14, 14, 14, 14, 16]
    thin = Side(style="thin", color="B7BBC5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, col, header); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill; cell.border = border; cell.alignment = center
        ws.column_dimensions[get_column_letter(col)].width = width
    for row, task in enumerate(tasks, 2):
        items = list(task.get("items") or [])
        def total(key: str) -> float | int:
            value = sum(float(item.get(key) or 0) for item in items)
            return int(value) if value.is_integer() else value
        values = [task.get("completedDate", ""), task.get("customer", ""), task.get("soNumber", ""),
                  task.get("containerType", ""), task.get("plannedShipDate", ""), task.get("port", ""),
                  len(items), total("quantity"), total("pieces"), total("volume")]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value); cell.border = border
            cell.alignment = left if col == 3 else center
        ws.row_dimensions[row].height = 30
    ws.auto_filter.ref = f"A1:J{max(2, ws.max_row)}"
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    date_range = f"{_text(payload.get('from'))}_{_text(payload.get('to'))}".strip("_") or date.today().isoformat()
    output = BytesIO(); wb.save(output)
    return output.getvalue(), f"{date_range}_走柜任务汇总.xlsx"
