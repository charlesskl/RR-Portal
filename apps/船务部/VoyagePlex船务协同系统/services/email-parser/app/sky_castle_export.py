"""Sky Castle single-cabinet notice: cabinet totals and PO context stay separate."""
from io import BytesIO
from decimal import Decimal, InvalidOperation
from .product_brand import product_brand
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.utils import get_column_letter


# Original customer template uses 0.0283 cubic metres per cubic foot.
CUBIC_METRES_PER_CUBIC_FOOT = Decimal("0.0283")
CUBIC_CENTIMETRES_PER_CUBIC_METRE = Decimal("1000000")


def calculated_cargo(source):
    item = dict(source)
    item["brand"] = product_brand(item)
    def number(key):
        try:
            value = Decimal(str(item.get(key)))
            return value if value.is_finite() and value > 0 else None
        except (InvalidOperation, ValueError): return None
    boxes = number("pieces"); spec = number("spec")
    if boxes is None and spec and number("quantity"):
        boxes = number("quantity") / spec
        if boxes == boxes.to_integral_value(): item["pieces"] = float(boxes)
        else: boxes = None
    if boxes and spec: item["quantity"] = float(boxes * spec)
    for total_key, unit_key in [("gross_weight", "gross_weight_per_box"), ("net_weight", "net_weight_per_box"), ("net_net_weight", "net_net_weight_per_box")]:
        if boxes and number(unit_key): item[total_key] = float(boxes * number(unit_key))
    try:
        dimensions = [Decimal(v.strip()) for v in str(item.get("box_dimensions") or "").replace("×", "*").split("*")]
        if len(dimensions) == 3 and all(v.is_finite() and v > 0 for v in dimensions):
            volume = dimensions[0] * dimensions[1] * dimensions[2] / CUBIC_CENTIMETRES_PER_CUBIC_METRE
            item["measurement_per_box"] = float(volume / CUBIC_METRES_PER_CUBIC_FOOT)
            if boxes: item["measurement"] = float(boxes * volume / CUBIC_METRES_PER_CUBIC_FOOT); item["volume"] = float(boxes * volume)
    except (InvalidOperation, ValueError): pass
    if boxes and number("measurement_per_box") and not number("measurement"):
        item["measurement"] = float(boxes * number("measurement_per_box"))
        item["volume"] = float(number("measurement") * CUBIC_METRES_PER_CUBIC_FOOT)
    elif number("volume") and not number("measurement"):
        item["measurement"] = float(number("volume") / CUBIC_METRES_PER_CUBIC_FOOT)
        if boxes: item["measurement_per_box"] = float(number("measurement") / boxes)
    return item


def build_sky_castle_workbook(task):
    from .shipment_export import _factory_note, _display_date, _display_deadline
    wb = Workbook(); ws = wb.active; ws.title = "走柜表"
    items = [calculated_cargo(item) for item in task.get("items") or []]
    end = max(14, 8 + len(items)); total = end + 1
    ws.merge_cells("I1:Q2"); ws["I1"] = "出货通知单"; ws["I1"].font = Font(name="宋体", size=20, bold=True)
    ws.merge_cells("A4:C4"); ws["A4"] = _display_date(task.get("plannedShipDate"))
    ws.merge_cells("E4:G4"); ws["E4"] = f"出{task.get('port') or ''}"
    ws.merge_cells("I4:T4"); ws["I4"] = f"RR {task.get('displayNumber') or task.get('id') or ''}柜"
    ws.merge_cells("A5:C5"); ws["A5"] = f"国名 {task.get('destinationCountry') or ''}"
    ws.merge_cells("D5:F5"); ws["D5"] = f"柜型 {task.get('containerType') or ''}"
    ws.merge_cells("G5:V5"); ws["G5"] = f"SO号码：{task.get('soNumber') or ''}"
    headers = ["洋行", "合同", "货号", "数量", "件数", "毛重", "净重", "净净重", "体积", "CBM", "玩具类别",
               "KG/箱（毛重）", "KG/箱（净重）", "KG/箱（净净重）", "尺码 CM", "", "", "每单总件数", "每箱/尺码", "拉", "客PO", "备注"]
    for col, label in enumerate(headers, 1): ws.cell(7, col, label)
    ws.merge_cells("O7:Q7")
    keys = [None, "contract_number", "product_code", "quantity", "pieces", "gross_weight", "net_weight", "net_net_weight",
            "measurement", "volume", "category", "gross_weight_per_box", "net_weight_per_box", "net_net_weight_per_box",
            None, None, None, "order_total_pieces", "measurement_per_box", "production_line", "customer_po", None]
    for row, item in enumerate(items, 8):
        for col, key in enumerate(keys, 1):
            value = item.get(key) if key else None
            if col == 1: value = "Sky Castle"
            if col == 3: value = "\n".join(str(v) for v in (item.get("product_code"), item.get("product_name")) if v)
            if col == 22: value = _factory_note(item, task.get("factoryMappings") or None)
            ws.cell(row, col, value if value != "" else None)
        dimensions = str(item.get("box_dimensions") or "").replace("×", "*").split("*")
        if len(dimensions) == 3:
            for col, value in enumerate(dimensions, 15):
                try: ws.cell(row, col, float(value))
                except ValueError: pass
        if item.get("quantity") and item.get("spec") and not item.get("pieces"):
            ws.cell(row, 5, float(item["quantity"]) / float(item["spec"]))
    ws.merge_cells(start_row=end, start_column=3, end_row=end, end_column=8)
    brands = list(dict.fromkeys(str(item["brand"]) for item in items if item.get("brand")))
    ws.cell(end, 3, "品牌：" + "、".join(brands)); ws.cell(end, 3).font = Font(name="宋体", size=18, color="FF0000", bold=True)
    ws.cell(total, 3, "合计")
    for col in range(4, 11):
        values = [ws.cell(row, col).value for row in range(8, 8 + len(items))]
        numbers = [Decimal(str(value)) for value in values if isinstance(value, (int, float))]
        ws.cell(total, col, float(sum(numbers)) if numbers else None)
    footer = [(f"截关期 {_display_deadline(task.get('cutoffDate'))}", "柜号："),
              (f"SI：{_display_deadline(task.get('siDeadline'))}", "封条："),
              ("收货人名称：", "柜重："), ("收货人代码：", ""),
              ("贸易方式：", "发货人："), ("拼箱：", "")]
    for offset, (left, right) in enumerate(footer, total+1):
        ws.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=5); ws.cell(offset, 1, left)
        ws.merge_cells(start_row=offset, start_column=6, end_row=offset, end_column=14); ws.cell(offset, 6, right)
    line = Side(style="thin", color="000000")
    for row in ws.iter_rows(min_row=7, max_row=end, max_col=22):
        for cell in row: cell.border = Border(left=line, right=line, top=line, bottom=line)
    for col in range(3, 11):
        ws.cell(total, col).border = Border(left=line, right=line, top=line, bottom=line)
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(horizontal="left" if cell.row >= end and cell.column != 3 else "center", vertical="center", wrap_text=True)
            if cell.font.name != "宋体": cell.font = Font(name="宋体", size=11)
    widths = [10.62,11.88,9.88,6.62,6.62,9.62,8.38,9,9,8.25,5.25,5.88,5.88,5.88,5.25,5.25,5.25,7,7.62,8.49,13.25,9]
    for col, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(8, total + 1):
        for col in [6,7,8,9,10,19]: ws.cell(row,col).number_format = "0.00"
        for col in [4,5,15,16,17,18]: ws.cell(row,col).number_format = "0"
    for row in [4,5]:
        ws.row_dimensions[row].height = 30
        for cell in ws[row]: cell.font = Font(name="宋体",size=18)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 14
    ws.row_dimensions[3].height = 4.5
    ws.row_dimensions[6].height = 6
    ws.row_dimensions[total].height = 24
    ws.cell(end,3).alignment = Alignment(horizontal="left",vertical="center")
    for row in range(total+1,total+7): ws.row_dimensions[row].height=20
    for row in [total+1,total+2,total+3]: ws.cell(row,6).font=Font(name="宋体",size=16,color="FF0000",bold=True)
    for row in range(8, end): ws.row_dimensions[row].height = 30
    ws.row_dimensions[7].height = 36
    for col in range(1,11): ws.cell(7,col).font = Font(name="宋体",size=14)
    for row in range(8,total+1):
        for col in range(4,11): ws.cell(row,col).font = Font(name="Times New Roman",size=12)
    for address in ("F7", "H7", "O7", "P7", "Q7"): ws[address].fill = PatternFill("solid", fgColor="FFFF00")
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.print_area = f"A1:V{total+6}"
    stream = BytesIO(); wb.save(stream)
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in str(task.get("soNumber") or task.get("id") or "待确认"))
    return stream.getvalue(), f"Sky_Castle_{safe}_走柜表.xlsx"
