import unittest
from io import BytesIO
from openpyxl import load_workbook
from app.factory_workbook import parse_factory_workbook
from openpyxl import Workbook
from app.factory_mapping import FACTORY_MAPPINGS, chinese_factory_name
from app.shipment_export import _factory_note, build_shipment_workbook


class FactoryMappingTests(unittest.TestCase):
    def test_cloud_records(self):
        self.assertEqual(len(FACTORY_MAPPINGS), 91)
        for row in FACTORY_MAPPINGS:
            self.assertEqual(chinese_factory_name(row["english_name"]), row["chinese_short_name"])
            self.assertIsInstance(row["is_local"], bool)

    def test_names_and_ambiguity(self):
        self.assertEqual(chinese_factory_name("Dong Guan Hanson Plastic Product Ltd"), "兴信")
        self.assertEqual(chinese_factory_name("  DANLI TOYS (LONGCHUAN) CO.,LTD. "), "丹尼")
        self.assertEqual(chinese_factory_name("Unknown Factory"), "Unknown Factory")
        self.assertEqual(chinese_factory_name("WINGMAU"), "WINGMAU")
        self.assertEqual(chinese_factory_name("WingMau"), "永贸")

    def test_export_actual_factory_and_preserve_source(self):
        item = {"supplier": "Dong Guan Hanson Plastic Product Ltd", "loading_factory": "华登"}
        self.assertEqual(_factory_note(item), "兴信")
        self.assertEqual(item["supplier"], "Dong Guan Hanson Plastic Product Ltd")

    def test_database_mapping_changes_export(self):
        records = [{"englishName": "Custom Factory", "chineseShortName": "新工厂", "isLocal": False}]
        data, _ = build_shipment_workbook({"company": "Xingxin", "customer": "ZURU", "factoryMappings": records, "items": [{"supplier": "Custom Factory"}]})
        self.assertEqual(load_workbook(BytesIO(data)).active["A6"].value, "新工厂")
        records[0]["chineseShortName"] = "已修改"
        data, _ = build_shipment_workbook({"customer": "ZURU", "factoryMappings": records, "items": [{"supplier": "Custom Factory"}]})
        self.assertEqual(load_workbook(BytesIO(data)).active["A6"].value, "已修改")

    def test_workbook_import_validation(self):
        wb = Workbook(); ws = wb.active
        ws.append(["英文名称", "中文简称", "本厂"])
        ws.append(["Factory", "工厂", "否"])
        content = BytesIO(); wb.save(content)
        self.assertEqual(parse_factory_workbook(content.getvalue()), [{"englishName": "Factory", "chineseShortName": "工厂", "isLocal": False}])
        ws["C2"] = "未知"; content = BytesIO(); wb.save(content)
        with self.assertRaises(ValueError): parse_factory_workbook(content.getvalue())

    def test_exported_remark_cells(self):
        data, _ = build_shipment_workbook({"company": "Xingxin", "customer": "ZURU", "items": [
            {"product_code": "77794", "supplier": "Dong Guan Hanson Plastic Product Ltd"},
            {"product_code": "15915", "supplier": "Danli Toys (Longchuan) Co.,Ltd"},
        ]})
        sheet = load_workbook(BytesIO(data)).active
        names = [row[0].value for row in sheet.iter_rows()]
        self.assertIn("兴信", names)
        self.assertIn("丹尼", names)
        self.assertNotIn("Dong Guan Hanson Plastic Product Ltd", names)


if __name__ == "__main__":
    unittest.main()
