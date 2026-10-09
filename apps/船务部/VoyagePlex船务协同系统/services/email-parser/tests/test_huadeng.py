from io import BytesIO
import os
import imaplib
import tempfile
from pathlib import Path
from email.message import EmailMessage
from app.eml import parse_eml
import unittest
from unittest.mock import patch, MagicMock
import openpyxl
from app.shipment_export import build_shipment_workbook as ordinary_builder
from app.huadeng_export import build_huadeng_workbook
def build_shipment_workbook(task): return build_huadeng_workbook(task, ordinary_builder)
from app.mailbox import fetch_mailbox, test_mailbox_connection
from app.main import scan_local_inventory

class HuadengTests(unittest.TestCase):
    def task(self):
        return {"company":"Huadeng","customer":"ZURU","plannedShipDate":"2026-08-29","soNumber":"22154703","port":"盐田","containerType":"1*40HQ","items":[{"product_code":"92107-S001","product_name":"摇篮公子8个/箱","quantity":7056,"pieces":882,"gross_weight":4207,"net_weight":3580,"volume":67.676,"category":"塑胶"}],"exportDetails":{"shipmentMode":"Container","sequenceNumber":"621","recipient":"余振军","containerNumber":"FANU1722982","sealNumber":"HLK4580036","containerWeight":"3680","consignee":"ROYAL REGENT TOYS PRODUCTS LTD."}}

    def test_container_layout_and_details(self):
        data, name=build_shipment_workbook(self.task())
        sheet=openpyxl.load_workbook(BytesIO(data)).active
        self.assertIn("华登",name)
        self.assertEqual("TO：余振军",sheet["A1"].value)
        self.assertIn("621柜",sheet["F2"].value)
        self.assertEqual("华登",sheet["A6"].value)
        self.assertEqual("塑胶",sheet["H6"].value)
        values=[str(c.value) for row in sheet for c in row if c.value is not None]
        self.assertTrue(any("东莞华登塑胶制品有限公司" in x for x in values))
        self.assertFalse(any("东莞兴信" in x for x in values))
        self.assertIn("柜号：FANU1722982",values)
        self.assertIn("船封：HLK4580036",values)
        self.assertEqual(3680,sheet["J10"].value)
        self.assertEqual("=K8+J9+J10",sheet["J11"].value)

    def test_warehouse_layout(self):
        task=self.task();task["exportDetails"].update(shipmentMode="Warehouse",sequenceNumber="510",warehouse="盐港二号仓",vehicleType="1*5T")
        data,name=build_shipment_workbook(task)
        sheet=openpyxl.load_workbook(BytesIO(data)).active
        self.assertIn("交仓走货表",name)
        self.assertIn("510车",sheet["F2"].value)
        self.assertIn("入盐港二号仓",sheet["E3"].value)
        self.assertIn("1*5T",sheet["E3"].value)
        self.assertFalse(any("船封：" in str(c.value) for row in sheet for c in row))

    def test_mixed_cartons_merge_and_count_once(self):
        task=self.task();task["items"]=[{"product_code":"MSLD200","product_name":"第一款","packing_group":"A","quantity":2058,"pieces":42,"gross_weight":2772,"net_weight":1337.7,"volume":60.627},{"product_code":"MSLD200","product_name":"第二款","packing_group":"A","quantity":882,"pieces":42,"gross_weight":2772,"net_weight":652.68,"volume":60.627}]
        data,_=build_shipment_workbook(task);sheet=openpyxl.load_workbook(BytesIO(data)).active
        self.assertIn("J6:J7",[str(x) for x in sheet.merged_cells.ranges])
        self.assertIn("K6:K7",[str(x) for x in sheet.merged_cells.ranges])
        self.assertEqual(42,sheet["J6"].value);self.assertIsNone(sheet["J7"].value)
        self.assertEqual(652.68,sheet["L7"].value)
        self.assertEqual("=SUM(J6:J7)",sheet["J9"].value)
        task["items"][1]["pieces"]=40
        with self.assertRaises(ValueError):build_shipment_workbook(task)

    def test_unconfigured_huadeng_never_uses_xingxin_mail_or_inventory(self):
        with patch.dict(os.environ,{"VOYAGEPLEX_MAIL_ADDRESS":"xingxin@example.com","VOYAGEPLEX_MAIL_AUTH_CODE":"secret","VOYAGEPLEX_LOCAL_INVENTORY_DIR":"/xingxin"},clear=True):
            self.assertFalse(fetch_mailbox(company="Huadeng")["configured"])
            self.assertIn("尚未配置",scan_local_inventory("Huadeng")["error"])

    def test_personal_connection_uses_supplied_credentials_and_folder(self):
        client=MagicMock()
        client.select.return_value=("OK", [])
        client.response.return_value=("UIDVALIDITY", [b"7"])
        client.uid.return_value=("OK", [b""])
        with patch.dict(os.environ,{"VOYAGEPLEX_HUADENG_MAIL_ADDRESS":"other@example.com","VOYAGEPLEX_HUADENG_MAIL_AUTH_CODE":"other"},clear=True), patch("app.mailbox.imaplib.IMAP4_SSL",return_value=client) as connection:
            result=fetch_mailbox(company="Huadeng",credentials={"address":"mine@example.com","auth_code":"personal-auth","host":"imap.example.com","folder":"Inbox/Shipping"})
            connection.assert_called_once_with("imap.example.com",993,timeout=30)
            client.login.assert_called_once_with("mine@example.com","personal-auth")
            client.select.assert_called_once_with("Inbox/Shipping",readonly=True)
            self.assertEqual("mine@example.com",result["address"])
            self.assertFalse(fetch_mailbox(company="Huadeng",credentials={})["configured"])

    def test_connection_check_never_searches_or_fetches_messages(self):
        client=MagicMock();client.select.return_value=("OK",[])
        credentials={"address":"mine@example.com","auth_code":"personal-auth","host":"imap.example.com","folder":"INBOX"}
        with patch("app.mailbox.imaplib.IMAP4_SSL",return_value=client):
            self.assertTrue(test_mailbox_connection(credentials)["connected"])
            client.uid.assert_not_called();client.select.assert_called_once_with("INBOX",readonly=True)
            client.logout.assert_called_once()
            client.login.side_effect=imaplib.IMAP4.error("personal-auth must never be returned")
            result=test_mailbox_connection(credentials)
            self.assertFalse(result["connected"]);self.assertNotIn("personal-auth",result["message"])
            client.login.side_effect=None;client.select.return_value=("NO",[])
            self.assertIn("文件夹",test_mailbox_connection(credentials)["message"])
        with patch("app.mailbox.imaplib.IMAP4_SSL",side_effect=OSError("secret")):
            self.assertIn("服务器",test_mailbox_connection(credentials)["message"])

    def test_same_notice_ignores_recipient_headers_but_checks_cargo_attachments(self):
        def message(recipient, content=b"cargo-a", body="安排装柜"):
            msg=EmailMessage();msg["From"]="sender@example.com";msg["To"]=recipient
            msg["Subject"]="走货通知";msg["Date"]="Thu, 08 Oct 2026 09:00:00 +0800"
            msg["Received"]="different delivery for "+recipient
            msg.set_content(body);msg.add_attachment(content,maintype="application",subtype="octet-stream",filename="cargo.xlsx")
            return msg.as_bytes()
        with tempfile.TemporaryDirectory() as directory:
            first=parse_eml(message("a@example.com"),Path(directory)/"a")
            second=parse_eml(message("b@example.com"),Path(directory)/"b")
            self.assertNotEqual(first["fingerprint"],second["fingerprint"])
            self.assertEqual(first["business_fingerprint"],second["business_fingerprint"])
            changed=parse_eml(message("b@example.com",content=b"cargo-b"),Path(directory)/"c")
            self.assertNotEqual(first["business_fingerprint"],changed["business_fingerprint"])
            changed=parse_eml(message("b@example.com",body="船期变更"),Path(directory)/"d")
            self.assertNotEqual(first["business_fingerprint"],changed["business_fingerprint"])

    def test_form_reads_merged_product_identity_without_duplicating_cartons(self):
        from app.huadeng_forms import parse_huadeng_form
        rows=[["备注","序号","洋行","合同","货号","货名","国家","玩具类别","数量","件数","毛重","净重","体积","客人PO#","每单总件数","每箱毛重","每箱净重"],
              ["华登",1,"ZURU","4500208073","MSLD200","第一款","美国","塑胶",2058,42,2772,1337.7,60.627,"PO",84,66,31.85],
              ["华登",2,"ZURU","","","第二款","美国","塑胶",882,"","",652.68,"","","","",15.54]]
        parsed=parse_huadeng_form(rows,[(1,3,4,5),(1,3,9,10)])
        self.assertEqual(2,len(parsed["items"]))
        self.assertEqual("MSLD200",parsed["items"][1]["product_code"])
        self.assertIsNone(parsed["items"][1]["pieces"])
        self.assertEqual(parsed["items"][0]["packing_group"],parsed["items"][1]["packing_group"])
