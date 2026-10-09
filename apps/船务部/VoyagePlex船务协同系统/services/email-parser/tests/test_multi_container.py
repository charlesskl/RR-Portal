from io import BytesIO
import unittest
from openpyxl import load_workbook
from app.carrier_schedule import ScheduleTables
from app.multi_container import multi_container_plan, order_cargo
from app.shipment_export import build_shipment_workbook
from app.product_workbook import parse_product_workbook


class MultiContainerTests(unittest.TestCase):
    def test_merged_schedule_rows_are_separate_orders(self):
        parser = ScheduleTables()
        parser.feed('<table><tr><th>BL#</th><th>SO#</th><th>Vessel name</th><th>SI cut</th><th>CY cut</th></tr>'
                    '<tr><td>BL111111</td><td>SO111111</td><td rowspan="2">VESSEL A</td><td rowspan="2">9/24(10:00)</td><td rowspan="2">9/27 (09:00 AM)</td></tr>'
                    '<tr><td>BL222222</td><td>SO222222</td></tr></table>')
        fields = {}; plan = multi_container_plan(fields, parser.tables, '2*40HQ', '2026-10-08')
        self.assertEqual(2, len(plan))
        self.assertEqual('SO222222', plan[1]['so_number'])
        self.assertEqual('VESSEL A', plan[1]['vessel_name'])
        self.assertEqual('2026-09-27T09:00', plan[1]['cutoff_date'])
        self.assertEqual([], multi_container_plan({}, parser.tables, '3*40HQ', '2026-10-08'))

    def test_order_table_extracts_only_provided_cargo(self):
        rows = [['Customer', 'Factory', 'Order Reference#', 'X3 PO#', 'Item#', 'Description', 'Qty(pc)', 'Qty/Ctn', 'No. of ctn', 'TTL CBM'],
                ['JAZWARES', 'STR', 'JAZ297155', 'JAZ297155', 'SR3703QP-12', 'STICKI ROLLS', '20016', '12', '1668', '437.54']]
        item = order_cargo([rows])[0]
        self.assertEqual(20016, item['quantity'])
        self.assertEqual(1668, item['pieces'])
        self.assertEqual('STICKI ROLLS', item['brand'])
        self.assertNotIn('net_net_weight', item)
        self.assertEqual([], order_cargo([[['CTN#', 'SEAL#']]]))

    def test_special_sheet_preserves_po_context_and_pure_product_weight(self):
        task = {'company':'Huadeng','soNumber':'SO111111','containerType':'40HQ','displayNumber':'56-1',
                'exportDetails':{'templateKey':'sky-castle-multi'},
                'items':[{'product_code':'SR3703QP-12','product_name':'手链','quantity':2868,'pieces':239,
                          'gross_weight':3393.8,'net_weight':2963.6,'net_net_weight':1118.52,
                          'gross_weight_per_box':14.2,'net_weight_per_box':12.4,'net_net_weight_per_box':4.68,
                          'box_dimensions':'59*36*123.5','order_total_pieces':1668,'brand':'STICKI ROLLS',
                          'supplier':'华登','category':'塑胶'}]}
        content, _ = build_shipment_workbook(task); sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(239, sheet['E8'].value)
        self.assertEqual(1668, sheet['R8'].value)
        self.assertEqual(1118.52, sheet['H8'].value)
        self.assertEqual('华登', sheet['V8'].value)
        self.assertEqual('FF0000', sheet['C14'].font.color.rgb[-6:])
        product = parse_product_workbook(content, 'special.xlsx')[0]
        self.assertEqual(12, product['quantityPerBox'])
        self.assertEqual(4.68, product['netNetWeightPerBox'])
        self.assertEqual('59*36*123.5', product['boxDimensions'])

    def test_export_calculates_values_and_matches_reference_layout(self):
        task = {'exportDetails':{'templateKey':'sky-castle-multi'},'items':[{
            'product_code':'SR3703QP-12','spec':12,'pieces':239,
            'gross_weight_per_box':14.2,'net_weight_per_box':12.4,'net_net_weight_per_box':4.68,
            'box_dimensions':'59*36*123.5','order_total_pieces':1668}]}
        content,_ = build_shipment_workbook(task)
        sheet=load_workbook(BytesIO(content),data_only=True).active
        for cell,expected in [('D8',2868),('F8',3393.8),('G8',2963.6),('H8',1118.52),('I8',2215.301978798587),('J8',62.693046),('S8',9.26904593639576)]:
            self.assertAlmostEqual(expected,sheet[cell].value,places=6)
        self.assertEqual(1668,sheet['R8'].value)
        self.assertIn('O7:Q7',str(sheet.merged_cells))
        self.assertEqual('0.00',sheet['J8'].number_format)
        self.assertEqual(3393.8,sheet['F15'].value)
        self.assertEqual('n',sheet['F8'].data_type)
        task['items'][0].pop('gross_weight_per_box')
        content,_=build_shipment_workbook(task)
        sheet=load_workbook(BytesIO(content),data_only=True).active
        self.assertIsNone(sheet['F8'].value)
        self.assertIsNone(sheet['F15'].value)

    def test_existing_product_brand_exports_without_reimport(self):
        task={'exportDetails':{'templateKey':'sky-castle-multi'},'items':[{'product_code':'SR3703QP-12','pieces':238,'spec':12}]}
        content,_=build_shipment_workbook(task)
        sheet=load_workbook(BytesIO(content)).active
        self.assertEqual('品牌：STICKI ROLLS',sheet['C14'].value)
        self.assertEqual('FF0000',sheet['C14'].font.color.rgb[-6:])
        task['items'][0]['brand']='人工确认品牌'
        content,_=build_shipment_workbook(task)
        self.assertEqual('品牌：人工确认品牌',load_workbook(BytesIO(content)).active['C14'].value)
        task['items'][0]={'product_code':'OTHER-SKU','pieces':238,'spec':12}
        content,_=build_shipment_workbook(task)
        self.assertEqual('品牌：',load_workbook(BytesIO(content)).active['C14'].value)

    def test_ordinary_huadeng_matches_existing_ordinary_template(self):
        task = {'company':'Huadeng','customer':'ZURU','soNumber':'SO111111','items':[]}
        first, _ = build_shipment_workbook(task)
        task['company'] = 'Xingxin'; second, _ = build_shipment_workbook(task)
        a = load_workbook(BytesIO(first)).active; b = load_workbook(BytesIO(second)).active
        self.assertEqual(list(a.values), list(b.values))
