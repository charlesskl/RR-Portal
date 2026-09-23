import unittest
from io import BytesIO

from openpyxl import Workbook

from app.product_workbook import parse_product_workbook


def workbook_bytes(headers, detail):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "柜单"
    for column, header in enumerate(headers, 1):
        sheet.cell(5, column, header)
    for column, value in enumerate(detail, 1):
        sheet.cell(6, column, value)
    sheet.cell(7, 2, "合计")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class ProductWorkbookTests(unittest.TestCase):
    def test_reads_detail_and_ignores_total_weights_and_footer(self):
        headers = ["备注", "序号", "客户", "合同", "货号", "货名 / 装箱规格", "国家", "玩具类别", "数量", "件数", "毛重", "净重", "体积", "客户PO#", "每单总件数", "每箱毛重", "每箱净重"]
        detail = ["兴信", 1, "ZURU", "C1", "77794-S001", "唱片机 16个/箱", "美国", "玩具", 160, 10, 50, 45, 1.2, "PO", 10, 5, 4.5]
        rows = parse_product_workbook(workbook_bytes(headers, detail), "走柜表.xlsx")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["productName"], "唱片机")
        self.assertEqual(rows[0]["quantityPerBox"], 16)
        self.assertEqual(rows[0]["grossWeightPerBox"], 5)
        self.assertEqual(rows[0]["netWeightPerBox"], 4.5)
        self.assertNotIn("factoryRemark", rows[0])

    def test_missing_spec_needs_review(self):
        headers = ["序号", "客户", "货号", "货名 / 装箱规格", "玩具类别", "每箱毛重", "每箱净重"]
        rows = parse_product_workbook(workbook_bytes(headers, [1, "TOMY", "A1", "特殊套装", "玩具", None, None]), "走柜表.xlsx")
        self.assertIsNone(rows[0]["quantityPerBox"])
        self.assertTrue(rows[0]["warnings"])


if __name__ == "__main__":
    unittest.main()
